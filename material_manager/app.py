"""工事材料管理システム — 統合スキーマ版"""
import os
import tempfile
from datetime import date

from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from openpyxl import load_workbook

from pdf_generator import generate_order_pdf
from pdf_parser import parse_order_pdf

from db import get_db, close_db, init_db, _now

app = Flask(__name__)
app.secret_key = "dev-key-change-in-production"
app.teardown_appcontext(close_db)


# ================================================================
# 現場一覧（project テーブル）
# ================================================================

@app.route("/")
def index():
    db = get_db()
    projects = db.execute(
        "SELECT * FROM project ORDER BY created_at DESC"
    ).fetchall()
    return render_template("index.html", sites=projects)


@app.route("/import-excel", methods=["POST"])
def import_excel():
    """Excelから現場名を読み取り、現場+見積もり明細を一括登録

    対応フォーマット:
    - 「鑑」シート E13 セルに工事件名
    - 「内訳書」シートに明細 (B=品目名, C=仕様, D=数量, E=単位, F=単価, G=金額)
    - ヘッダ行あり、ページ区切り行あり → 自動スキップ
    """
    db = get_db()

    file = request.files.get("excel_file")
    if not file or not file.filename.endswith(('.xlsx', '.xls')):
        flash("Excelファイル(.xlsx)を選択してください。")
        return redirect(url_for("index"))

    # パラメータ取得
    site_name_cell = request.form.get("site_name_cell", "E13").strip()
    site_name_sheet = request.form.get("site_name_sheet", "鑑").strip()
    detail_sheet = request.form.get("detail_sheet", "内訳書").strip()
    category = request.form.get("category", "")
    col_name = int(request.form.get("col_name", 2)) - 1    # B列
    col_spec = int(request.form.get("col_spec", 3)) - 1    # C列
    col_qty = int(request.form.get("col_qty", 4)) - 1      # D列
    col_unit = int(request.form.get("col_unit", 5)) - 1    # E列
    col_price = int(request.form.get("col_price", 6)) - 1  # F列
    col_amount = int(request.form.get("col_amount", 7)) - 1 # G列
    start_row = int(request.form.get("start_row", 3))

    # ファイル保存して読み込み
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        file.save(tmp.name)
        tmp_path = tmp.name

    try:
        wb = load_workbook(tmp_path, read_only=True, data_only=True)

        # 現場名を「鑑」シートから読み取り
        site_name = None
        for sname in wb.sheetnames:
            if site_name_sheet in sname:
                ws_cover = wb[sname]
                site_name = ws_cover[site_name_cell].value
                break

        if not site_name or str(site_name).strip() == "":
            # ファイル名から抽出
            site_name = file.filename.replace(".xlsx", "").replace(".xls", "")
            # "見積書_K2509-001_" のプレフィクスを除去
            parts = site_name.split("_")
            if len(parts) >= 3:
                site_name = "_".join(parts[2:])
        site_name = str(site_name).strip()

        # 現場を登録
        cur = db.execute(
            """INSERT INTO project (title, category, status, created_at, updated_at)
               VALUES (?, ?, '施工中', ?, ?)""",
            (site_name, category, _now(), _now())
        )
        db.commit()
        project_id = cur.lastrowid

        # 見積もりヘッダ作成
        cur = db.execute(
            "INSERT INTO estimate_header (project_id, version, created_at) VALUES (?, 1, ?)",
            (project_id, _now())
        )
        db.commit()
        header_id = cur.lastrowid

        # 「内訳書」シートから明細読み込み
        ws = None
        for sname in wb.sheetnames:
            if detail_sheet in sname:
                ws = wb[sname]
                break
        if ws is None:
            ws = wb.active

        # スキップすべき行のパターン
        skip_keywords = ['品', '小', '合', '計', '労務費', '現場雑費', 'その他']

        imported = 0
        for row in ws.iter_rows(min_row=start_row, values_only=True):
            name = row[col_name] if col_name < len(row) else None
            if not name or str(name).strip() == "":
                continue

            name = str(name).strip()

            # ヘッダ行・小計行・合計行をスキップ
            if name.startswith('品') and '名' in name:
                continue
            if any(name.startswith(k) for k in ['  小', '  合']):
                continue

            # 数量チェック
            try:
                qty = float(row[col_qty]) if col_qty < len(row) and row[col_qty] else 0
            except (ValueError, TypeError):
                continue
            if qty == 0:
                continue

            # 単位
            unit = str(row[col_unit]).strip() if col_unit < len(row) and row[col_unit] else "式"

            # 単価
            try:
                price = int(float(row[col_price])) if col_price < len(row) and row[col_price] else 0
            except (ValueError, TypeError):
                price = 0

            # 金額（金額列優先）
            amount = int(qty * price) if price else 0
            if col_amount < len(row) and row[col_amount]:
                try:
                    amount = int(float(row[col_amount]))
                except (ValueError, TypeError):
                    pass
            if amount == 0:
                continue

            # 仕様をスペックとして取得
            spec = str(row[col_spec]).strip() if col_spec < len(row) and row[col_spec] else ""

            # 品目名+仕様で識別（同名品目でも仕様違いを区別）
            full_name = f"{name} {spec}".strip() if spec else name

            # 品目マスタ（既存なら再利用）
            existing_item = db.execute(
                "SELECT id FROM item_master WHERE name = ? AND spec = ?", (name, spec)
            ).fetchone()
            if existing_item:
                item_id = existing_item["id"]
            else:
                c = db.execute(
                    "INSERT INTO item_master (name, spec, unit, category, standard_price, updated_at) "
                    "VALUES (?, ?, ?, ?, ?, ?)",
                    (name, spec, unit, category, price, _now())
                )
                db.commit()
                item_id = c.lastrowid

            db.execute(
                "INSERT INTO estimate_line (header_id, item_id, quantity, unit_price, amount, sort_order) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (header_id, item_id, qty, price, amount, imported)
            )
            imported += 1

        # ヘッダ合計更新
        total = db.execute(
            "SELECT COALESCE(SUM(amount), 0) FROM estimate_line WHERE header_id = ?",
            (header_id,)
        ).fetchone()[0]
        db.execute("UPDATE estimate_header SET total_amount = ? WHERE id = ?", (total, header_id))
        db.commit()
        wb.close()

        flash(f"「{site_name}」を登録し、{imported} 件の見積もり明細をインポートしました。")
    except Exception as e:
        flash(f"エラー: {e}")
    finally:
        os.unlink(tmp_path)

    return redirect(url_for("index"))


@app.route("/site/<int:project_id>/edit", methods=["POST"])
def site_edit(project_id):
    """現場情報の編集（担当者名・住所・発注元・工事種別）"""
    db = get_db()
    db.execute(
        """UPDATE project SET manager = ?, client = ?, address = ?, category = ?, updated_at = ?
           WHERE id = ?""",
        (request.form.get("manager", ""), request.form.get("client", ""),
         request.form.get("address", ""), request.form.get("category", ""),
         _now(), project_id)
    )
    db.commit()
    flash("現場情報を更新しました。")
    return redirect(url_for("site_detail", project_id=project_id))


@app.route("/site/new", methods=["GET", "POST"])
def site_new():
    if request.method == "POST":
        db = get_db()
        db.execute(
            """INSERT INTO project (title, address, manager, category, status, created_at, updated_at)
               VALUES (?, ?, ?, ?, '施工中', ?, ?)""",
            (request.form["name"], request.form["address"],
             request.form["manager"], request.form.get("category", ""),
             _now(), _now())
        )
        db.commit()
        flash("現場を登録しました。")
        return redirect(url_for("index"))
    return render_template("site_form.html")


# ================================================================
# 現場詳細: 見積もり vs 発注比較
# ================================================================

@app.route("/site/<int:project_id>")
def site_detail(project_id):
    db = get_db()
    site = db.execute("SELECT * FROM project WHERE id = ?", (project_id,)).fetchone()

    # 最新版の見積もりヘッダを取得（なければ None）
    header = db.execute(
        "SELECT * FROM estimate_header WHERE project_id = ? ORDER BY version DESC LIMIT 1",
        (project_id,)
    ).fetchone()

    rows = []
    total_est = 0
    total_order = 0

    if header:
        rows = db.execute("""
            SELECT
                im.id AS item_id,
                im.name AS item_name,
                im.spec,
                im.unit,
                el.quantity AS est_qty,
                el.unit_price AS est_unit_price,
                el.amount AS est_amount,
                COALESCE(o.order_qty, 0) AS order_qty,
                COALESCE(o.order_amount, 0) AS order_amount,
                COALESCE(rc.missing_receipts, 0) AS missing_receipts
            FROM estimate_line el
            JOIN item_master im ON im.id = el.item_id
            LEFT JOIN (
                SELECT item_id,
                       SUM(quantity) AS order_qty,
                       SUM(amount) AS order_amount
                FROM "order"
                WHERE project_id = ?
                GROUP BY item_id
            ) o ON o.item_id = el.item_id
            LEFT JOIN (
                SELECT ord.item_id,
                       COUNT(*) - COUNT(r.id) AS missing_receipts
                FROM "order" ord
                LEFT JOIN receipt r ON r.order_id = ord.id
                WHERE ord.project_id = ?
                GROUP BY ord.item_id
            ) rc ON rc.item_id = el.item_id
            WHERE el.header_id = ?
            ORDER BY el.sort_order, im.name
        """, (project_id, project_id, header["id"])).fetchall()

        total_est = sum(r["est_amount"] for r in rows)
        total_order = sum(r["order_amount"] for r in rows)

    return render_template("site_detail.html",
                           site=site, rows=rows,
                           total_est=total_est, total_order=total_order)


# ================================================================
# 見積もり Excel インポート
# ================================================================

@app.route("/site/<int:project_id>/estimate/import", methods=["GET", "POST"])
def estimate_import(project_id):
    """Excelファイルから見積もり明細を一括インポート"""
    db = get_db()
    site = db.execute("SELECT * FROM project WHERE id = ?", (project_id,)).fetchone()

    if request.method == "POST":
        file = request.files.get("excel_file")
        if not file or not file.filename.endswith(('.xlsx', '.xls')):
            flash("Excelファイル(.xlsx)を選択してください。")
            return redirect(request.url)

        # 列マッピング（ユーザー指定）
        col_name = int(request.form.get("col_name", 1)) - 1
        col_unit = int(request.form.get("col_unit", 3)) - 1
        col_qty = int(request.form.get("col_qty", 4)) - 1
        col_price = int(request.form.get("col_price", 5)) - 1
        col_amount = int(request.form.get("col_amount", 6)) - 1
        start_row = int(request.form.get("start_row", 2))
        category = request.form.get("category", "")

        # ファイル保存して読み込み
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            file.save(tmp.name)
            tmp_path = tmp.name

        try:
            wb = load_workbook(tmp_path, read_only=True, data_only=True)
            ws = wb.active

            # 見積もりヘッダ作成（既存なら新版）
            existing = db.execute(
                "SELECT MAX(version) as v FROM estimate_header WHERE project_id = ?",
                (project_id,)
            ).fetchone()
            new_version = (existing["v"] or 0) + 1

            cur = db.execute(
                "INSERT INTO estimate_header (project_id, version, created_by, created_at) VALUES (?, ?, ?, ?)",
                (project_id, new_version, request.form.get("created_by", ""), _now())
            )
            db.commit()
            header_id = cur.lastrowid

            imported = 0
            for row_idx, row in enumerate(ws.iter_rows(min_row=start_row, values_only=True), start=start_row):
                # 品目名が空ならスキップ
                name = row[col_name] if col_name < len(row) else None
                if not name or str(name).strip() == "":
                    continue

                name = str(name).strip()
                unit = str(row[col_unit]).strip() if col_unit < len(row) and row[col_unit] else "式"

                # 数量・単価を取得
                try:
                    qty = float(row[col_qty]) if col_qty < len(row) and row[col_qty] else 0
                except (ValueError, TypeError):
                    continue
                if qty == 0:
                    continue

                try:
                    price = int(float(row[col_price])) if col_price < len(row) and row[col_price] else 0
                except (ValueError, TypeError):
                    price = 0

                amount = int(qty * price)
                # 金額列があればそちらを優先
                if col_amount < len(row) and row[col_amount]:
                    try:
                        amount = int(float(row[col_amount]))
                    except (ValueError, TypeError):
                        pass

                # 品目マスタに登録（既存なら再利用）
                existing_item = db.execute(
                    "SELECT id FROM item_master WHERE name = ?", (name,)
                ).fetchone()
                if existing_item:
                    item_id = existing_item["id"]
                else:
                    c = db.execute(
                        "INSERT INTO item_master (name, unit, category, updated_at) VALUES (?, ?, ?, ?)",
                        (name, unit, category, _now())
                    )
                    db.commit()
                    item_id = c.lastrowid

                # 明細行追加
                db.execute(
                    "INSERT INTO estimate_line (header_id, item_id, quantity, unit_price, amount, sort_order) "
                    "VALUES (?, ?, ?, ?, ?, ?)",
                    (header_id, item_id, qty, price, amount, imported)
                )
                imported += 1

            # ヘッダの合計更新
            total = db.execute(
                "SELECT COALESCE(SUM(amount), 0) FROM estimate_line WHERE header_id = ?",
                (header_id,)
            ).fetchone()[0]
            db.execute("UPDATE estimate_header SET total_amount = ? WHERE id = ?", (total, header_id))
            db.commit()
            wb.close()

            flash(f"Excel から {imported} 件の見積もり明細をインポートしました（版 {new_version}）。")
        except Exception as e:
            flash(f"エラー: {e}")
        finally:
            os.unlink(tmp_path)

        return redirect(url_for("site_detail", project_id=project_id))

    return render_template("estimate_import.html", site=site)


# ================================================================
# 見積もり明細追加（手動）
# ================================================================

@app.route("/site/<int:project_id>/estimate/add", methods=["GET", "POST"])
def estimate_add(project_id):
    db = get_db()
    site = db.execute("SELECT * FROM project WHERE id = ?", (project_id,)).fetchone()

    if request.method == "POST":
        # 品目: 既存 or 新規
        item_id = request.form.get("item_id")
        if not item_id or item_id == "new":
            cur = db.execute(
                "INSERT INTO item_master (name, unit, category, updated_at) VALUES (?, ?, ?, ?)",
                (request.form["new_item_name"], request.form["new_item_unit"],
                 request.form.get("new_item_category", ""), _now())
            )
            db.commit()
            item_id = cur.lastrowid
        else:
            item_id = int(item_id)

        qty = float(request.form["quantity"])
        unit_price = int(float(request.form["unit_price"]))
        amount = int(qty * unit_price)

        # 見積もりヘッダ: この案件の最新版を取得、なければ新規作成
        header = db.execute(
            "SELECT * FROM estimate_header WHERE project_id = ? ORDER BY version DESC LIMIT 1",
            (project_id,)
        ).fetchone()

        if not header:
            cur = db.execute(
                "INSERT INTO estimate_header (project_id, version, created_by, created_at) VALUES (?, 1, ?, ?)",
                (project_id, request.form.get("created_by", ""), _now())
            )
            db.commit()
            header_id = cur.lastrowid
        else:
            header_id = header["id"]

        # 明細行追加
        db.execute(
            "INSERT INTO estimate_line (header_id, item_id, quantity, unit_price, amount, memo) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (header_id, item_id, qty, unit_price, amount, request.form.get("note", ""))
        )

        # ヘッダの合計を更新
        total = db.execute(
            "SELECT COALESCE(SUM(amount), 0) FROM estimate_line WHERE header_id = ?",
            (header_id,)
        ).fetchone()[0]
        db.execute(
            "UPDATE estimate_header SET total_amount = ? WHERE id = ?",
            (total, header_id)
        )
        db.commit()

        # 単価履歴に記録（AI学習用）
        db.execute(
            "INSERT INTO price_history (item_id, unit_price, recorded_date, source) VALUES (?, ?, ?, '見積もり')",
            (item_id, unit_price, date.today().isoformat())
        )
        db.commit()

        flash("見積もり明細を追加しました。")
        return redirect(url_for("site_detail", project_id=project_id))

    items = db.execute("SELECT * FROM item_master ORDER BY category, name").fetchall()
    return render_template("estimate_form.html", site=site, items=items)


# ================================================================
# 発注登録
# ================================================================

@app.route("/site/<int:project_id>/order/add", methods=["GET", "POST"])
def order_add(project_id):
    db = get_db()
    site = db.execute("SELECT * FROM project WHERE id = ?", (project_id,)).fetchone()

    if request.method == "POST":
        item_id = int(request.form["item_id"])
        qty = float(request.form["quantity"])
        unit_price = int(float(request.form["unit_price"]))
        amount = int(qty * unit_price)

        # 仕入先: 既存 or 新規登録
        supplier_name = request.form.get("supplier", "").strip()
        supplier_id = None
        if supplier_name:
            existing = db.execute(
                "SELECT id FROM supplier WHERE name = ?", (supplier_name,)
            ).fetchone()
            if existing:
                supplier_id = existing["id"]
            else:
                cur = db.execute(
                    "INSERT INTO supplier (name, created_at) VALUES (?, ?)",
                    (supplier_name, _now())
                )
                db.commit()
                supplier_id = cur.lastrowid

        db.execute(
            """INSERT INTO "order" (project_id, item_id, supplier_id, quantity, unit_price,
               amount, order_date, orderer, status, memo) VALUES (?, ?, ?, ?, ?, ?, ?, ?, '発注済', ?)""",
            (project_id, item_id, supplier_id, qty, unit_price, amount,
             request.form["order_date"], request.form["orderer"],
             request.form.get("note", ""))
        )

        # 単価履歴に記録（AI学習用）
        db.execute(
            "INSERT INTO price_history (item_id, supplier_id, unit_price, recorded_date, source) "
            "VALUES (?, ?, ?, ?, '発注')",
            (item_id, supplier_id, unit_price, request.form["order_date"])
        )
        db.commit()

        flash("発注を登録しました。")
        return redirect(url_for("site_detail", project_id=project_id))

    # この案件の見積もり品目
    header = db.execute(
        "SELECT id FROM estimate_header WHERE project_id = ? ORDER BY version DESC LIMIT 1",
        (project_id,)
    ).fetchone()

    if header:
        items = db.execute("""
            SELECT DISTINCT im.* FROM item_master im
            JOIN estimate_line el ON el.item_id = im.id
            WHERE el.header_id = ?
            ORDER BY im.name
        """, (header["id"],)).fetchall()
    else:
        items = db.execute("SELECT * FROM item_master ORDER BY name").fetchall()

    return render_template("order_form.html", site=site, items=items,
                           today=date.today().isoformat())


# ================================================================
# 品目ドリルダウン（発注履歴）
# ================================================================

@app.route("/site/<int:project_id>/item/<int:item_id>")
def item_orders(project_id, item_id):
    db = get_db()
    site = db.execute("SELECT * FROM project WHERE id = ?", (project_id,)).fetchone()
    item = db.execute("SELECT * FROM item_master WHERE id = ?", (item_id,)).fetchone()

    # 見積もり情報
    header = db.execute(
        "SELECT id FROM estimate_header WHERE project_id = ? ORDER BY version DESC LIMIT 1",
        (project_id,)
    ).fetchone()

    estimate = None
    if header:
        estimate = db.execute(
            "SELECT * FROM estimate_line WHERE header_id = ? AND item_id = ?",
            (header["id"], item_id)
        ).fetchone()

    # 発注履歴（仕入先名 + 受領書有無を結合）
    orders = db.execute("""
        SELECT o.*, COALESCE(s.name, '') AS supplier_name,
               (SELECT COUNT(*) FROM receipt r WHERE r.order_id = o.id) AS receipt_count
        FROM "order" o
        LEFT JOIN supplier s ON s.id = o.supplier_id
        WHERE o.project_id = ? AND o.item_id = ?
        ORDER BY o.order_date
    """, (project_id, item_id)).fetchall()

    total_qty = sum(o["quantity"] for o in orders)
    total_amount = sum(o["amount"] for o in orders)
    has_missing_receipt = any(o["receipt_count"] == 0 for o in orders)

    return render_template("item_orders.html",
                           site=site, item=item, estimate=estimate,
                           orders=orders, total_qty=total_qty,
                           total_amount=total_amount,
                           has_missing_receipt=has_missing_receipt)


# ================================================================
# 受領書アップロード
# ================================================================

@app.route("/order/<int:order_id>/receipt/upload", methods=["POST"])
def receipt_upload(order_id):
    """発注に対する受領書をアップロード"""
    db = get_db()
    order = db.execute('SELECT * FROM "order" WHERE id = ?', (order_id,)).fetchone()
    if not order:
        flash("発注レコードが見つかりません。")
        return redirect(url_for("index"))

    file = request.files.get("receipt_file")
    if not file or file.filename == "":
        flash("ファイルを選択してください。")
        return redirect(request.referrer or url_for("index"))

    # ファイル保存
    import uuid
    ext = os.path.splitext(file.filename)[1]
    saved_name = f"receipt_{order_id}_{uuid.uuid4().hex[:8]}{ext}"
    save_dir = os.path.join(app.root_path, "uploads", "receipts")
    os.makedirs(save_dir, exist_ok=True)
    save_path = os.path.join(save_dir, saved_name)
    file.save(save_path)

    uploaded_by = request.form.get("uploaded_by", "")

    db.execute(
        "INSERT INTO receipt (order_id, file_path, file_name, uploaded_by, uploaded_at) VALUES (?, ?, ?, ?, ?)",
        (order_id, saved_name, file.filename, uploaded_by, _now())
    )
    db.commit()

    flash(f"受領書をアップロードしました。（{file.filename}）")
    return redirect(request.referrer or url_for("site_detail", project_id=order["project_id"]))


@app.route("/receipts/<path:filename>")
def receipt_file(filename):
    """受領書ファイルのダウンロード"""
    return send_file(os.path.join(app.root_path, "uploads", "receipts", filename))


# ================================================================
# 発注書 PDF インポート（PDFから発注データを読み取り登録）
# ================================================================

@app.route("/site/<int:project_id>/order/import-pdf", methods=["GET", "POST"])
def order_import_pdf(project_id):
    """Step 1: PDFアップロード → パース → プレビュー表示"""
    db = get_db()
    site = db.execute("SELECT * FROM project WHERE id = ?", (project_id,)).fetchone()

    if request.method == "POST":
        file = request.files.get("pdf_file")
        if not file or not file.filename.endswith('.pdf'):
            flash("PDFファイルを選択してください。")
            return redirect(request.url)

        col_map = {
            "name": int(request.form.get("col_name", 1)) - 1,
            "spec": int(request.form.get("col_spec", 2)) - 1 if request.form.get("col_spec") else None,
            "qty": int(request.form.get("col_qty", 3)) - 1,
            "unit": int(request.form.get("col_unit", 4)) - 1,
            "price": int(request.form.get("col_price", 5)) - 1,
            "amount": int(request.form.get("col_amount", 6)) - 1,
        }
        start_row = int(request.form.get("start_row", 1))
        orderer = request.form.get("orderer", "")
        supplier_name = request.form.get("supplier", "")
        order_date = request.form.get("order_date", date.today().isoformat())

        # PDF保存してパース
        upload_dir = os.path.join(app.root_path, "uploads")
        os.makedirs(upload_dir, exist_ok=True)
        tmp_path = os.path.join(upload_dir, f"tmp_{project_id}_{date.today().isoformat()}.pdf")
        file.save(tmp_path)

        try:
            items, metadata = parse_order_pdf(tmp_path, col_map=col_map, start_row=start_row)
        except Exception as e:
            flash(f"PDF解析エラー: {e}")
            os.unlink(tmp_path)
            return redirect(request.url)

        if not orderer and metadata.get("orderer"):
            orderer = metadata["orderer"]
        if not supplier_name and metadata.get("supplier"):
            supplier_name = metadata["supplier"]
        if not order_date and metadata.get("order_date"):
            order_date = metadata["order_date"]

        # 各品目のマッチング判定
        existing_items = db.execute("SELECT id, name, spec, unit FROM item_master ORDER BY name").fetchall()
        preview = []
        for i, item_data in enumerate(items):
            name = item_data["name"]
            spec = item_data.get("spec", "")

            # マッチング: 名前+仕様 → 名前のみ → 部分一致
            match_id = None
            match_type = "unmatched"
            for ei in existing_items:
                if ei["name"] == name and ei["spec"] == spec:
                    match_id = ei["id"]
                    match_type = "exact"
                    break
            if not match_id:
                for ei in existing_items:
                    if ei["name"] == name:
                        match_id = ei["id"]
                        match_type = "name_only"
                        break
            if not match_id:
                for ei in existing_items:
                    if name in ei["name"] or ei["name"] in name:
                        match_id = ei["id"]
                        match_type = "partial"
                        break

            preview.append({
                "idx": i,
                "name": name,
                "spec": spec,
                "qty": item_data["qty"],
                "unit": item_data["unit"],
                "unit_price": item_data["unit_price"],
                "amount": item_data["amount"],
                "match_id": match_id,
                "match_type": match_type,
            })

        return render_template("order_import_preview.html",
                               site=site, preview=preview,
                               existing_items=existing_items,
                               orderer=orderer, supplier=supplier_name,
                               order_date=order_date, tmp_path=tmp_path)

    return render_template("order_import_pdf.html", site=site, today=date.today().isoformat())


@app.route("/site/<int:project_id>/order/import-pdf/confirm", methods=["POST"])
def order_import_pdf_confirm(project_id):
    """Step 2: ユーザーの振り分け確認後に実際に登録"""
    db = get_db()
    orderer = request.form.get("orderer", "")
    supplier_name = request.form.get("supplier", "")
    order_date = request.form.get("order_date", date.today().isoformat())
    tmp_path = request.form.get("tmp_path", "")

    # 仕入先登録
    supplier_id = None
    if supplier_name:
        existing = db.execute("SELECT id FROM supplier WHERE name = ?", (supplier_name,)).fetchone()
        if existing:
            supplier_id = existing["id"]
        else:
            cur = db.execute("INSERT INTO supplier (name, created_at) VALUES (?, ?)",
                             (supplier_name, _now()))
            db.commit()
            supplier_id = cur.lastrowid

    # 各品目を処理
    idx = 0
    imported = 0
    while True:
        name = request.form.get(f"item_{idx}_name")
        if name is None:
            break

        action = request.form.get(f"item_{idx}_action", "skip")
        if action == "skip":
            idx += 1
            continue

        qty = float(request.form.get(f"item_{idx}_qty", 0))
        unit = request.form.get(f"item_{idx}_unit", "式")
        unit_price = int(float(request.form.get(f"item_{idx}_price", 0)))
        amount = int(float(request.form.get(f"item_{idx}_amount", 0)))
        spec = request.form.get(f"item_{idx}_spec", "")

        if action == "new":
            # 新規品目として登録
            cur = db.execute(
                "INSERT INTO item_master (name, spec, unit, category, standard_price, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (name, spec, unit, "", unit_price, _now())
            )
            db.commit()
            item_id = cur.lastrowid
        else:
            # 既存品目に振り分け
            item_id = int(action)

        # 発注レコード登録
        db.execute(
            """INSERT INTO "order" (project_id, item_id, supplier_id, quantity, unit_price,
               amount, order_date, orderer, status, memo) VALUES (?, ?, ?, ?, ?, ?, ?, ?, '発注済', '')""",
            (project_id, item_id, supplier_id, qty, unit_price, amount, order_date, orderer)
        )

        # 単価履歴
        db.execute(
            "INSERT INTO price_history (item_id, supplier_id, unit_price, recorded_date, source) "
            "VALUES (?, ?, ?, ?, '発注PDF')",
            (item_id, supplier_id, unit_price, order_date)
        )
        imported += 1
        idx += 1

    db.commit()

    # 一時ファイル削除
    if tmp_path and os.path.exists(tmp_path):
        os.unlink(tmp_path)

    flash(f"発注書PDFから {imported} 件の発注データを登録しました。")
    return redirect(url_for("site_detail", project_id=project_id))


# ================================================================
# 発注書 PDF 生成
# ================================================================

@app.route("/site/<int:project_id>/item/<int:item_id>/pdf", methods=["GET", "POST"])
def item_order_pdf(project_id, item_id):
    """品目ごとの発注書PDF生成"""
    db = get_db()
    site = db.execute("SELECT * FROM project WHERE id = ?", (project_id,)).fetchone()
    item = db.execute("SELECT * FROM item_master WHERE id = ?", (item_id,)).fetchone()

    if request.method == "POST":
        # 発注履歴を取得
        orders = db.execute("""
            SELECT o.*, COALESCE(s.name, '') AS supplier_name
            FROM "order" o
            LEFT JOIN supplier s ON s.id = o.supplier_id
            WHERE o.project_id = ? AND o.item_id = ?
            ORDER BY o.order_date
        """, (project_id, item_id)).fetchall()

        order_list = [dict(o) for o in orders]

        pdf_buf = generate_order_pdf(
            site_title=site["title"],
            item_name=item["name"],
            item_spec=item["spec"] or "",
            item_unit=item["unit"],
            orders=order_list,
            supplier_name=request.form.get("supplier_name", ""),
            orderer_name=request.form.get("orderer_name", ""),
            delivery_date=request.form.get("delivery_date", ""),
            delivery_place=request.form.get("delivery_place", ""),
            note=request.form.get("note", ""),
        )

        filename = f"発注書_{item['name']}_{date.today().isoformat()}.pdf"
        return send_file(pdf_buf, mimetype="application/pdf",
                         as_attachment=True, download_name=filename)

    return render_template("order_pdf_form.html", site=site, item=item)


# ================================================================
# CLI
# ================================================================

@app.cli.command("init-db")
def init_db_command():
    init_db()
    print("DB を初期化しました。")


if __name__ == "__main__":
    app.run(debug=True, port=5000)
