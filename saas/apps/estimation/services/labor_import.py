"""社内整形済み Excel からの労務単価取込サービス。

注意: 国土交通省は公共工事設計労務単価を **PDFのみ** で配布しており、
公式Excelは存在しない。公式PDFからの取込は services/labor_pdf.py を使うこと。

このモジュールは、社内で手整形したExcel（神奈川県分など少数職種を
先行投入する運用）の取込口として使う。
"""

from datetime import date
from decimal import Decimal, InvalidOperation

import openpyxl

from apps.estimation.models import LaborRate

# 47都道府県のリスト（マッチング用）
PREFECTURES = [
    "北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県",
    "茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県",
    "新潟県", "富山県", "石川県", "福井県", "山梨県", "長野県",
    "岐阜県", "静岡県", "愛知県", "三重県",
    "滋賀県", "京都府", "大阪府", "兵庫県", "奈良県", "和歌山県",
    "鳥取県", "島根県", "岡山県", "広島県", "山口県",
    "徳島県", "香川県", "愛媛県", "高知県",
    "福岡県", "佐賀県", "長崎県", "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県",
]

# 都道府県名のバリエーション（"県"なし等）
_PREF_LOOKUP = {}
for p in PREFECTURES:
    _PREF_LOOKUP[p] = p
    if p.endswith(("県", "都", "府")):
        _PREF_LOOKUP[p[:-1]] = p
    elif p == "北海道":
        _PREF_LOOKUP["北海道"] = p


def _normalize_prefecture(text: str) -> str | None:
    """都道府県名を正規化する。マッチしなければ None。"""
    if not text:
        return None
    text = str(text).strip()
    return _PREF_LOOKUP.get(text)


def _is_number(value) -> bool:
    """セル値が数値として扱えるか。"""
    if value is None:
        return False
    try:
        Decimal(str(value))
        return True
    except (InvalidOperation, ValueError):
        return False


def _detect_header_row(ws):
    """ヘッダー行（職種名が並んでいる行）を検出する。"""
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=False), 1):
        for cell in row:
            if cell.value and _normalize_prefecture(str(cell.value)):
                return row_idx - 1, cell.column - 1
    return None, None


def parse_labor_rate_excel(file_path: str) -> list[dict]:
    """社内整形済み Excel を解析し、レコードのリストを返す。

    Returns:
        [{"prefecture": "神奈川県", "trade": "電工", "amount": 25600}, ...]
        未設定職種（空セルや非数値セル）はスキップされる。
    """
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    records = []

    for ws in wb.worksheets:
        header_row_idx, pref_col_idx = _detect_header_row(ws)
        if header_row_idx is None:
            continue

        header_row = list(ws.iter_rows(
            min_row=header_row_idx, max_row=header_row_idx, values_only=True,
        ))[0]

        trades = {}
        for col_idx, cell_value in enumerate(header_row):
            if col_idx == pref_col_idx:
                continue
            if cell_value and str(cell_value).strip():
                trade_name = str(cell_value).strip()
                if len(trade_name) >= 2 and not trade_name.isdigit():
                    trades[col_idx] = trade_name

        if not trades:
            continue

        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not row or pref_col_idx >= len(row):
                continue
            pref = _normalize_prefecture(str(row[pref_col_idx]) if row[pref_col_idx] else "")
            if not pref:
                continue

            for col_idx, trade_name in trades.items():
                if col_idx >= len(row):
                    continue
                value = row[col_idx]
                if not _is_number(value):
                    # 未設定職種: スキップ（NULLとして扱う）
                    continue
                amount = int(Decimal(str(value)))
                if amount <= 0:
                    continue
                records.append({
                    "prefecture": pref,
                    "trade": trade_name,
                    "amount": amount,
                })

    wb.close()
    return records


def import_labor_rates_from_excel(
    file_path: str,
    valid_from,
    company,
    *,
    source_label: str = "",
    fiscal_year_label: str = "",
) -> dict:
    """社内整形済みExcelから労務単価を取り込む。

    注意: 国土交通省はPDFのみ配布しており、公式Excelは存在しない。
    公式PDFからの取込は services/labor_pdf.py を使うこと。

    Returns:
        {"created": int, "updated": int, "total": int, "batch_id": str}
    """
    batch_id = f"excel_{valid_from}_{date.today().isoformat()}"
    records = parse_labor_rate_excel(file_path)

    created = 0
    updated = 0

    for rec in records:
        _, was_created = LaborRate.unscoped.update_or_create(  # unscoped: company を明示指定
            company=company,
            prefecture=rec["prefecture"],
            occupation_code="",
            valid_from=valid_from,
            defaults={
                "occupation_name": rec["trade"],
                "unit_price": Decimal(str(rec["amount"])),
                "fiscal_year_label": fiscal_year_label,
                "status": "draft",
                "import_batch": batch_id,
            },
        )
        if was_created:
            created += 1
        else:
            updated += 1

    return {
        "created": created,
        "updated": updated,
        "total": len(records),
        "batch_id": batch_id,
    }
