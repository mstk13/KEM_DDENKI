"""全省庁統一資格 省庁別許可内容一覧: Excel 読み込み・登録コマンド・一覧表示・越境。"""

import pytest
from django.core.management import call_command

from apps.bids.importer import import_unified_excel
from apps.bids.models import UnifiedQualification
from apps.core.tenant_context import set_current_company


@pytest.fixture
def unified_xlsx(tmp_path):
    """実物と同じ列構成（結合セル・2段見出し・注記行）の Excel を作る。"""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "基本情報"
    ws = wb.create_sheet("省庁別許可内容一覧")
    ws["A1"] = "省庁別 許可区分・等級・点数・営業品目(できること)一覧"
    ws["A2"] = "※注記"
    header = {
        "A4": "No.", "B4": "省庁・機関名", "C4": "押印", "D4": "押印者(会計担当部署)",
        "E4": "備考(所属・区分)", "F4": "物品の製造", "H4": "物品の販売",
        "J4": "物品の販売:営業品目(できること)", "K4": "役務の提供等",
        "M4": "役務の提供等:営業品目(できること)", "N4": "物品の買受け",
        "P4": "物品の買受け:営業品目",
    }
    for k, v in header.items():
        ws[k] = v
    for k in ("F5", "H5", "K5", "N5"):
        ws[k] = "等級"
    for k in ("G5", "I5", "L5", "O5"):
        ws[k] = "点数"
    for rng in ("F4:G4", "H4:I4", "K4:L4", "N4:O4"):
        ws.merge_cells(rng)
    rows = [
        (1, "衆議院", "○", "庶務部会計課長", "立法機関", "―", "―",
         "C", 62, "電気・通信用機器類/その他", "C", 62, "賃貸借/その他", "B", 62, "その他"),
        (2, "防衛省", "○", "大臣官房会計課長", None, "―", "―",
         "C", 62, "電気・通信用機器類/その他", "C", 62, "賃貸借/その他", "B", 62, "その他"),
    ]
    for i, row in enumerate(rows, start=6):
        for j, val in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=val)
    ws["A9"] = "※「物品の製造」は資格登録なし"
    path = tmp_path / "unified.xlsx"
    wb.save(path)
    return path


class TestImportUnifiedExcel:
    def test_reads_agency_rows_only(self, unified_xlsx):
        records = import_unified_excel(unified_xlsx)
        assert [r["agency"] for r in records] == ["衆議院", "防衛省"]
        assert [r["sort_order"] for r in records] == [1, 2]

    def test_maps_grade_score_and_items(self, unified_xlsx):
        rec = import_unified_excel(unified_xlsx)[0]
        assert rec["goods_sales_grade"] == "C"
        assert rec["goods_sales_score"] == 62
        assert rec["goods_sales_items"] == "電気・通信用機器類/その他"
        assert rec["services_grade"] == "C"
        assert rec["services_score"] == 62
        assert rec["services_items"] == "賃貸借/その他"
        assert rec["purchase_grade"] == "B"
        assert rec["purchase_score"] == 62
        # 「物品の製造」は取り込まない
        assert "manufacturing" not in " ".join(rec.keys())

    def test_dash_means_blank(self, tmp_path, unified_xlsx):
        import openpyxl

        wb = openpyxl.load_workbook(unified_xlsx)
        ws = wb["省庁別許可内容一覧"]
        ws["N6"] = "―"
        ws["O6"] = "―"
        path = tmp_path / "dash.xlsx"
        wb.save(path)
        rec = import_unified_excel(path)[0]
        assert rec["purchase_grade"] == ""
        assert rec["purchase_score"] is None

    def test_missing_header_raises(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        wb.create_sheet("省庁別許可内容一覧")["A1"] = "見出しなし"
        path = tmp_path / "bad.xlsx"
        wb.save(path)
        with pytest.raises(ValueError):
            import_unified_excel(path)


@pytest.mark.django_db
class TestImportCommand:
    def test_creates_rows_for_first_company(self, company_a, unified_xlsx):
        call_command("import_unified_qualifications", str(unified_xlsx))
        qs = UnifiedQualification.unscoped.filter(company=company_a).order_by("sort_order")
        assert [u.agency for u in qs] == ["衆議院", "防衛省"]

    def test_rerun_updates_instead_of_duplicating(self, company_a, unified_xlsx):
        call_command("import_unified_qualifications", str(unified_xlsx))
        call_command("import_unified_qualifications", str(unified_xlsx))
        assert UnifiedQualification.unscoped.filter(company=company_a).count() == 2

    def test_replace_removes_old_rows(self, company_a, user_a, unified_xlsx):
        UnifiedQualification.unscoped.create(
            company=company_a, created_by=user_a, agency="旧機関", sort_order=99,
        )
        call_command("import_unified_qualifications", str(unified_xlsx), "--replace")
        agencies = set(
            UnifiedQualification.unscoped.filter(company=company_a)
            .values_list("agency", flat=True)
        )
        assert agencies == {"衆議院", "防衛省"}


@pytest.mark.django_db
class TestUnifiedQualificationIsolation:
    def test_isolation(self, company_a, company_b, user_a, user_b):
        set_current_company(company_a)
        UnifiedQualification.objects.create(
            company=company_a, created_by=user_a, agency="衆議院",
        )
        assert UnifiedQualification.objects.count() == 1

        set_current_company(company_b)
        assert UnifiedQualification.objects.count() == 0
        set_current_company(None)

    def test_other_tenant_cannot_edit(self, client, company_a, user_a, user_b):
        obj = UnifiedQualification.unscoped.create(
            company=company_a, created_by=user_a, agency="衆議院",
        )
        client.force_login(user_b)
        response = client.get(f"/bids/qualifications/unified/{obj.pk}/edit/")
        assert response.status_code == 404


@pytest.mark.django_db
class TestUnifiedQualificationViews:
    def test_list_shows_second_table(self, client, company_a, user_a):
        UnifiedQualification.unscoped.create(
            company=company_a, created_by=user_a, agency="衆議院",
            goods_sales_grade="C", goods_sales_score=62,
            goods_sales_items="電気・通信用機器類/その他",
            services_grade="C", services_score=62, services_items="賃貸借/その他",
            purchase_grade="B", purchase_score=62,
        )
        client.force_login(user_a)
        response = client.get("/bids/qualifications/")
        assert response.status_code == 200
        html = response.content.decode()
        assert "全省庁統一資格 省庁別許可内容一覧" in html
        assert "衆議院" in html
        assert "電気・通信用機器類/その他" in html

    def test_create_edit_delete(self, client, company_a, user_a):
        client.force_login(user_a)
        response = client.post("/bids/qualifications/unified/new/", {
            "sort_order": 1, "agency": "総務省",
            "goods_sales_grade": "C", "goods_sales_score": 62, "goods_sales_items": "その他",
            "services_grade": "C", "services_score": 62, "services_items": "その他",
            "purchase_grade": "B", "purchase_score": 62,
        })
        assert response.status_code == 302
        obj = UnifiedQualification.unscoped.get(company=company_a, agency="総務省")
        assert obj.created_by == user_a

        response = client.post(f"/bids/qualifications/unified/{obj.pk}/edit/", {
            "sort_order": 1, "agency": "総務省", "goods_sales_grade": "B",
            "goods_sales_score": 70, "goods_sales_items": "その他",
            "services_grade": "C", "services_score": 62, "services_items": "その他",
            "purchase_grade": "B", "purchase_score": 62,
        })
        assert response.status_code == 302
        obj.refresh_from_db()
        assert obj.goods_sales_grade == "B"
        assert obj.goods_sales_score == 70

        response = client.post(f"/bids/qualifications/unified/{obj.pk}/delete/")
        assert response.status_code == 302
        assert not UnifiedQualification.unscoped.filter(pk=obj.pk).exists()
