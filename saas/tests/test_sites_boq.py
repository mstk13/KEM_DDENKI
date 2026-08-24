"""現場の内訳書・内訳明細書のテスト。

- モデル制約（持ち主は積算案件か現場のどちらか一方）
- テナント越境
- Excel / PDF の取込パーサ
- 階層の組み立て（並び順と階層から親子を作る）
- 現場詳細での表示、表形式編集、Excel出力
"""

import io
from decimal import Decimal
from types import SimpleNamespace

import openpyxl
import pytest
from django.db import IntegrityError, transaction
from django.urls import reverse

from apps.estimation.models import BoqLine, EstimationProject, Orderer
from apps.estimation.services import boq_import
from apps.estimation.services.boq_export import export_boq_to_excel
from apps.sites.models import Site


def _site(company, code="S-001", name="○○ビル電気設備工事"):
    return Site.unscoped.create(company=company, code=code, name=name)


def _line(company, site, name, level=BoqLine.Level.SAIMOKU, order=1, **kwargs):
    return BoqLine.unscoped.create(
        company=company, site=site, name=name, level=level,
        sort_order=order, **kwargs,
    )


def _excel_bytes(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADER = ["階層", "名称", "仕様", "単位", "数量", "単価", "金額", "備考"]


# ---------------------------------------------------------------------------
# モデル
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestBoqLineOwner:
    def test_site_only_is_allowed(self, company_a):
        site = _site(company_a)
        line = _line(company_a, site, "VVFケーブル")
        assert line.site_id == site.pk
        assert line.project_id is None

    def test_neither_owner_is_rejected(self, company_a):
        """持ち主のいない明細は作れない。"""
        with pytest.raises(IntegrityError), transaction.atomic():
            BoqLine.unscoped.create(
                company=company_a, name="宙に浮いた明細",
                level=BoqLine.Level.SAIMOKU,
            )

    def test_both_owners_are_rejected(self, company_a):
        """両方に紐づけるとどちらの内訳書に出すか決まらない。"""
        site = _site(company_a)
        orderer = Orderer.unscoped.create(company=company_a, name="○○市")
        project = EstimationProject.unscoped.create(
            company=company_a, name="案件", orderer=orderer,
        )
        with pytest.raises(IntegrityError), transaction.atomic():
            BoqLine.unscoped.create(
                company=company_a, site=site, project=project,
                name="二重", level=BoqLine.Level.SAIMOKU,
            )

    def test_is_meisai_only_for_saimoku(self, company_a):
        site = _site(company_a)
        assert _line(company_a, site, "細目", BoqLine.Level.SAIMOKU).is_meisai is True
        assert _line(company_a, site, "科目", BoqLine.Level.KAMOKU).is_meisai is False

    def test_calc_amount_keeps_imported_amount(self, company_a):
        """一式計上の行の金額を計算で消さない。"""
        site = _site(company_a)
        line = _line(
            company_a, site, "仮設工事一式",
            quantity=None, unit_price=None, amount=Decimal("500000"),
        )
        line.calc_amount()
        assert line.amount == Decimal("500000")


@pytest.mark.django_db
class TestBoqTenantIsolation:
    def test_lines_do_not_cross_tenants(self, company_a, company_b):
        from apps.core.tenant_context import set_current_company

        site_a = _site(company_a, "A-001")
        site_b = _site(company_b, "B-001")
        _line(company_a, site_a, "A社の明細")
        _line(company_b, site_b, "B社の明細")

        set_current_company(company_a)
        names = list(BoqLine.objects.values_list("name", flat=True))
        assert names == ["A社の明細"]

    def test_other_tenant_site_is_not_reachable(self, client, user_a, company_b):
        site_b = _site(company_b, "B-001")
        client.force_login(user_a)
        response = client.get(reverse("sites:boq_edit", args=[site_b.pk]))
        assert response.status_code == 404


# ---------------------------------------------------------------------------
# 値の正規化
# ---------------------------------------------------------------------------


class TestValueParsing:
    @pytest.mark.parametrize(("raw", "expected"), [
        ("1,234", Decimal("1234")),
        ("¥1,234", Decimal("1234")),
        ("1234円", Decimal("1234")),
        ("１２３４", Decimal("1234")),      # 全角
        ("12.500", Decimal("12.500")),
        ("-500", Decimal("-500")),
        (1234, Decimal("1234")),
        (12.5, Decimal("12.5")),
    ])
    def test_numbers(self, raw, expected):
        assert boq_import._to_decimal(raw) == expected

    @pytest.mark.parametrize("raw", ["", None, "一式", "式", "-", "―", "　"])
    def test_unreadable_becomes_none(self, raw):
        """読めない値は None。0 にすると読めた 0 と区別できなくなる。"""
        assert boq_import._to_decimal(raw) is None

    def test_lump_sum_with_number_is_read(self):
        """「1式」でも数字があるものは読む。"""
        assert boq_import._to_decimal("1式") == Decimal("1")

    @pytest.mark.parametrize("name", ["小計", "合計", "計", "小　計", "総計"])
    def test_subtotal_rows_are_detected(self, name):
        assert boq_import._is_subtotal(name) is True

    def test_normal_name_is_not_subtotal(self):
        assert boq_import._is_subtotal("電灯設備工事") is False


class TestIndentDepth:
    @pytest.mark.parametrize(("raw", "depth"), [
        ("電気設備工事", 0),
        ("  電灯設備", 1),
        ("　電灯設備", 1),        # 全角空白1つ
        ("    配線器具", 2),
        ("　　配線器具", 2),
        ("      VVFケーブル", 3),
    ])
    def test_depth(self, raw, depth):
        assert boq_import._indent_depth(raw) == depth


# ---------------------------------------------------------------------------
# Excel 取込
# ---------------------------------------------------------------------------


class TestExcelImport:
    def test_reads_rows_with_level_column(self):
        data = _excel_bytes([
            ["工事名称", "○○ビル電気設備工事"],
            [],
            HEADER,
            ["種目別", "電気設備工事", "", "式", None, None, 5000000, ""],
            ["細目別", "VVFケーブル", "1.6mm 2芯", "m", 1200, 150, 180000, "屋内"],
        ])
        result = boq_import.parse_excel(data)

        assert result.warnings == []
        assert len(result.drafts) == 2
        assert result.drafts[0].level == BoqLine.Level.SHUMOKU
        assert result.drafts[1].level == BoqLine.Level.SAIMOKU
        assert result.drafts[1].spec == "1.6mm 2芯"
        assert result.drafts[1].quantity == Decimal("1200")
        assert result.drafts[1].unit_price == Decimal("150")
        assert result.drafts[1].amount == Decimal("180000")
        assert result.meisai_count == 1

    def test_infers_level_from_indent(self):
        """全角空白1つで1段。4階層が素直に読める。"""
        data = _excel_bytes([
            ["名称", "仕様", "単位", "数量", "単価", "金額"],
            ["電気設備工事", "", "式", None, None, 5000000],
            ["　電灯設備", "", "式", None, None, 3000000],
            ["　　配線工事", "", "式", None, None, 1000000],
            ["　　　VVFケーブル", "1.6-2C", "m", 1200, 150, 180000],
        ])
        result = boq_import.parse_excel(data)

        levels = [d.level for d in result.drafts]
        assert levels == [
            BoqLine.Level.SHUMOKU,
            BoqLine.Level.KAMOKU,
            BoqLine.Level.CHUKAMOKU,
            BoqLine.Level.SAIMOKU,
        ]

    def test_priced_row_is_meisai_regardless_of_indent(self):
        """数量×単価がそろう行は、インデントが浅くても内訳明細書。

        実ファイルのインデントは2段だったり4段だったりばらつくので、
        段数だけで見ると最下層が細目別にならない。内訳明細書とは
        「数量×単価で金額を積む階層」のことなので、そこで判定する。
        """
        data = _excel_bytes([
            ["名称", "規格・仕様", "単位", "数量", "単価", "金額"],
            ["電気設備工事", "", "式", None, None, 5180000],
            ["　電灯設備", "", "式", None, None, 180000],
            ["　　VVFケーブル", "1.6mm 2芯", "m", 1200, 150, 180000],
        ])
        result = boq_import.parse_excel(data)

        levels = [d.level for d in result.drafts]
        assert levels == [
            BoqLine.Level.SHUMOKU,
            BoqLine.Level.KAMOKU,
            BoqLine.Level.SAIMOKU,
        ]
        assert result.meisai_count == 1

    def test_lump_sum_row_is_not_meisai(self):
        """一式計上（金額だけ）の行は見出し扱いのまま。"""
        data = _excel_bytes([
            ["名称", "単位", "数量", "単価", "金額"],
            ["　仮設工事", "式", "一式", None, 500000],
        ])
        result = boq_import.parse_excel(data)

        assert result.drafts[0].quantity is None
        assert result.drafts[0].level != BoqLine.MEISAI_LEVEL

    def test_top_level_row_with_amount_is_not_demoted(self):
        """インデントのある表では、最上位行が金額を持っていても種目別のまま。

        インデントの有無を行ごとに見ていた頃、金額を持つ最上位行だけが
        数値の有無で判定され細目別になっていた。
        """
        data = _excel_bytes([
            ["名称", "単位", "数量", "単価", "金額"],
            ["電気設備工事", "式", None, None, 5000000],
            ["　VVFケーブル", "m", 1200, 150, 180000],
        ])
        result = boq_import.parse_excel(data)

        assert result.drafts[0].level == BoqLine.Level.SHUMOKU

    def test_infers_level_from_numbers_when_flat(self):
        """インデントも階層列も無い様式。数量があれば明細とみなす。"""
        data = _excel_bytes([
            ["名称", "単位", "数量", "単価", "金額"],
            ["電灯設備", "", None, None, None],
            ["VVFケーブル", "m", 1200, 150, 180000],
        ])
        result = boq_import.parse_excel(data)

        assert result.drafts[0].level == BoqLine.Level.KAMOKU
        assert result.drafts[1].level == BoqLine.Level.SAIMOKU

    def test_skips_subtotal_rows(self):
        data = _excel_bytes([
            HEADER,
            ["細目別", "VVFケーブル", "", "m", 1200, 150, 180000, ""],
            ["", "小計", "", "", None, None, 180000, ""],
            ["", "合計", "", "", None, None, 180000, ""],
        ])
        result = boq_import.parse_excel(data)

        assert [d.name for d in result.drafts] == ["VVFケーブル"]

    def test_column_alias_is_absorbed(self):
        data = _excel_bytes([
            ["品名", "規格", "単位", "数量", "単価", "金額"],
            ["VVFケーブル", "1.6-2C", "m", 1200, 150, 180000],
        ])
        result = boq_import.parse_excel(data)

        assert result.drafts[0].name == "VVFケーブル"
        assert result.drafts[0].spec == "1.6-2C"

    def test_cover_sheet_is_skipped_without_noise(self):
        """表紙しかないシートで警告を出さない（本当の失敗が埋もれる）。"""
        wb = openpyxl.Workbook()
        cover = wb.active
        cover.title = "鑑"
        cover.append(["工事名称", "○○ビル"])
        body = wb.create_sheet("内訳書")
        for row in [HEADER, ["細目別", "VVF", "", "m", 10, 100, 1000, ""]]:
            body.append(row)
        buf = io.BytesIO()
        wb.save(buf)

        result = boq_import.parse_excel(buf.getvalue())
        assert result.warnings == []
        assert len(result.drafts) == 1

    def test_unreadable_file_reports_reason(self):
        result = boq_import.parse_excel(b"not an excel file")
        assert result.drafts == []
        assert result.warnings
        assert "Excel を開けません" in result.warnings[0]

    def test_no_table_reports_reason(self):
        data = _excel_bytes([["工事名称", "○○ビル"], ["発注者", "△△市"]])
        result = boq_import.parse_excel(data)
        assert result.drafts == []
        assert result.warnings


class TestParseUploadDispatch:
    def test_xls_is_rejected_with_guidance(self):
        result = boq_import.parse_upload("内訳書.xls", b"")
        assert ".xlsx で保存し直して" in result.warnings[0]

    def test_unknown_extension_is_rejected(self):
        result = boq_import.parse_upload("内訳書.docx", b"")
        assert result.drafts == []
        assert result.warnings


# ---------------------------------------------------------------------------
# PDF 取込
# ---------------------------------------------------------------------------
#
# 実物の見積書PDFを持ち込まずに、実ファイルで踏んだ壊れ方を再現する。
#
# - 縞模様の背景を敷いた PDF は罫線が1行おきにしか引かれず、
#   `extract_tables()` が罫線の無い行の全列を1セルに潰す
# - 内訳書と内訳明細書が同じPDFに綴じられている
# - 単位列を持たず、数量セルに「0.66㎥(立方メートル)」と入る
# - 階層は名称の左端位置で示され、種目行にも「数量1・単価・金額」が入る

TABLE_TOP = 100.0
ROW_PITCH = 20.0
# 縦罫線の位置。No / 名称 / 仕様 / 数量 / 単価 / 金額 / 備考 の7列。
COLUMN_RULES = [30.0, 60.0, 260.0, 450.0, 530.0, 610.0, 690.0, 810.0]
# 名称列の左端。段の深さをここで表す様式。
NAME_INDENTS = [65.0, 75.0, 85.0, 95.0]


def _word(text, x0, top):
    return {
        "text": text, "x0": x0, "x1": x0 + max(len(text) * 5.0, 5.0),
        "top": top, "bottom": top + 10.0,
    }


class _FakePage:
    """pdfplumber の Page のうち、取込が触る部分だけを持つ差し替え。

    `rows` は (段, No, 名称, 仕様, 数量, 単価, 金額) のならび。
    """

    HEADER = ["No", "名称", "仕様", "数量", "単価", "金額", "備考"]

    def __init__(self, title, rows, *, page_number=1,
                 footer="※印は軽減税率対象です"):
        self.title = title
        self.rows = rows
        self.page_number = page_number
        self.footer = footer
        self.height = 600.0

    # -- 座標 ---------------------------------------------------------------

    @property
    def _bottom(self):
        return TABLE_TOP + (len(self.rows) + 1) * ROW_PITCH

    @property
    def edges(self):
        edges = [
            {"orientation": "v", "x0": x, "x1": x,
             "top": TABLE_TOP, "bottom": self._bottom}
            for x in COLUMN_RULES
        ]
        edges += [
            {"orientation": "h", "x0": COLUMN_RULES[0], "x1": COLUMN_RULES[-1],
             "top": y, "bottom": y}
            for y in (TABLE_TOP, self._bottom)
        ]
        return edges

    def extract_words(self):
        words = [
            _word(label, COLUMN_RULES[i] + 5.0, TABLE_TOP + 4.0)
            for i, label in enumerate(self.HEADER)
        ]
        for index, row in enumerate(self.rows):
            depth, no, name, spec, qty, price, amount = row
            top = TABLE_TOP + (index + 1) * ROW_PITCH + 4.0
            if no:
                words.append(_word(no, COLUMN_RULES[0] + 5.0, top))
            words.append(_word(name, NAME_INDENTS[depth], top))
            for column, value in ((2, spec), (3, qty), (4, price), (5, amount)):
                if value:
                    words.append(_word(value, COLUMN_RULES[column] + 5.0, top))
        # 表の外の脚注。名称列の最左になって段の順位をずらしていた。
        words.append(_word(self.footer, COLUMN_RULES[0], self._bottom + 30.0))
        return words

    # -- pdfplumber の API --------------------------------------------------

    def extract_text(self):
        lines = [self.title, "○○工事 一式", " ".join(self.HEADER)]
        lines += [" ".join(str(v) for v in row[1:] if v) for row in self.rows]
        lines.append(self.footer)
        return "\n".join(lines)

    def extract_tables(self):
        """罫線が1行おきにしか引かれない表を返す。

        奇数行は7列ぶんが1セルに潰れ、名称列が空になる。
        """
        table = [list(self.HEADER)]
        for index, row in enumerate(self.rows):
            cells = [row[1], row[2], row[3], row[4], row[5], row[6], ""]
            if index % 2:
                jammed = " ".join(c for c in cells if c)
                table.append([jammed, "", "", "", "", "", ""])
            else:
                table.append(cells)
        return [table]


class _FakePdf:
    def __init__(self, pages):
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


@pytest.fixture
def fake_pdf(monkeypatch):
    def install(pages):
        for number, page in enumerate(pages, start=1):
            page.page_number = number
        monkeypatch.setattr(boq_import, "HAS_PDFPLUMBER", True)
        monkeypatch.setattr(
            boq_import, "pdfplumber",
            SimpleNamespace(open=lambda _data: _FakePdf(pages)),
        )
        return pages
    return install


SUMMARY_ROWS = [
    (0, "A", "舗装工", "", "1", "¥ 70,000", "¥ 70,000"),
    (0, "B", "掘削工事", "", "1", "¥ 96,080", "¥ 96,080"),
]
DETAIL_ROWS_1 = [
    (0, "A", "舗装工", "", "1", "¥ 70,000", "¥ 70,000"),
    (1, "a", "アスファルト撤去工", "", "1", "¥ 70,000", "¥ 70,000"),
    (2, "1", "アスファルトカッター工", "", "1", "¥ 26,000", "¥ 26,000"),
    (3, "1", "施工費", "約10.4m/カッター工", "10.4m", "¥ 2,500", "¥ 26,000"),
]
DETAIL_ROWS_2 = [
    (2, "2", "アスファルト撤去", "", "1", "¥ 44,000", "¥ 44,000"),
    (3, "1", "施工費", "約6.4m²", "6.4㎡", "¥ 6,875", "¥ 44,000"),
    (0, "B", "掘削工事", "", "1", "¥ 96,080", "¥ 96,080"),
    (1, "a", "掘削工", "", "1", "¥ 96,080", "¥ 96,080"),
    (2, "1", "施工費", "2.88㎥", "2.88㎥(", "¥ 33,361", "¥ 96,080"),
    (0, "", "【合計】", "", "", "", "¥ 166,080"),
]


def _ninomiya_pages():
    return [
        _FakePage("内訳書", SUMMARY_ROWS),
        _FakePage("内訳明細書", DETAIL_ROWS_1),
        _FakePage("内訳明細書", DETAIL_ROWS_2),
    ]


class TestQuantityWithUnitInSameCell:
    """単位列を持たない様式。数量セルに単位が同居する。"""

    @pytest.mark.parametrize(("raw", "quantity", "unit"), [
        ("0.66㎡", Decimal("0.66"), "㎡"),
        ("2.88㎥", Decimal("2.88"), "㎥"),
        ("6.4㎡", Decimal("6.4"), "㎡"),
        ("0.66㎥(\n立方メートル)", Decimal("0.66"), "㎥"),
        ("10.4m", Decimal("10.4"), "m"),
        ("8本", Decimal("8"), "本"),
        ("1式", Decimal("1"), "式"),
    ])
    def test_unit_does_not_leak_into_the_number(self, raw, quantity, unit):
        """NFKC は ㎡ を "m2"、㎥ を "m3" に展開する。

        数字以外を消す実装では、その "2" "3" が数量の末尾に残り
        「0.66㎡」が 0.662 になっていた。金額は ¥ しか付かないので
        気づかれず、数量×単価だけが合わなくなる。
        """
        assert boq_import._to_decimal(raw) == quantity
        assert boq_import._split_unit(raw) == unit

    def test_unit_column_wins_when_present(self):
        data = _excel_bytes([
            ["名称", "単位", "数量", "単価", "金額"],
            ["VVFケーブル", "m", 1200, 150, 180000],
        ])
        result = boq_import.parse_excel(data)
        assert result.drafts[0].unit == "m"


class TestSubtotalBrackets:
    @pytest.mark.parametrize("name", ["【合計】", "【小計】", "（合計）", "[計]"])
    def test_bracketed_subtotal_is_detected(self, name):
        """「【合計】」と括る様式がある。取り込むと二重計上になる。"""
        assert boq_import._is_subtotal(name) is True


class TestPageKind:
    @pytest.mark.parametrize(("title", "kind"), [
        ("内訳明細書", "detail"),
        ("細目別内訳書", "detail"),
        ("内訳書", "summary"),
        ("種目別内訳書", "summary"),
        ("御見積書", None),
    ])
    def test_kind_from_heading(self, title, kind):
        assert boq_import._page_kind(_FakePage(title, SUMMARY_ROWS)) == kind


class TestRowsFromWords:
    def test_every_row_is_recovered(self):
        """罫線が1行おきの表でも、語の座標から全行を組み直せる。"""
        page = _FakePage("内訳明細書", DETAIL_ROWS_1)

        collapsed = page.extract_tables()[0]
        assert boq_import._dropped_row_count(collapsed) == 2

        rows = boq_import._rows_from_words(page)
        assert boq_import._usable_row_count(rows) == len(DETAIL_ROWS_1)

    def test_footer_outside_the_table_is_ignored(self):
        """脚注を拾うと名称列の最左になり、段の順位が1つずつずれる。"""
        page = _FakePage("内訳明細書", DETAIL_ROWS_1)
        rows = boq_import._rows_from_words(page)

        assert all("※印" not in "".join(row) for row in rows)

    def test_name_indent_is_restored(self):
        """名称列の左端を全角空白に置き換え、Excel と同じ判定に乗せる。"""
        page = _FakePage("内訳明細書", DETAIL_ROWS_1)
        rows = boq_import._rows_from_words(page)
        header_index, mapping = boq_import._find_header(rows)
        names = [row[mapping["name"]] for row in rows[header_index + 1:]]

        assert [boq_import._indent_depth(n) for n in names] == [0, 1, 2, 3]


class TestPdfImport:
    def test_striped_pdf_keeps_every_row(self, fake_pdf):
        """罫線が1行おきでも行が落ちない。落ちた行は名称が空で捨てられていた。"""
        fake_pdf(_ninomiya_pages())
        result = boq_import.parse_pdf(b"")

        assert result.warnings == []
        # 【合計】の1行だけが小計行として落ちる。
        assert len(result.drafts) == len(DETAIL_ROWS_1) + len(DETAIL_ROWS_2) - 1

    def test_summary_pages_are_skipped(self, fake_pdf):
        """内訳書は内訳明細書の集計。両方取り込むと二重計上になる。"""
        fake_pdf(_ninomiya_pages())
        result = boq_import.parse_pdf(b"")

        shumoku = [d for d in result.drafts if d.level == BoqLine.Level.SHUMOKU]
        assert [d.name for d in shumoku] == ["舗装工", "掘削工事"]
        assert sum(d.amount for d in shumoku) == Decimal("166080")

    def test_summary_only_pdf_is_still_read(self, fake_pdf):
        """内訳明細書が無いなら内訳書を読む。飛ばすのは重複するときだけ。"""
        fake_pdf([_FakePage("内訳書", SUMMARY_ROWS)])
        result = boq_import.parse_pdf(b"")

        assert [d.name for d in result.drafts] == ["舗装工", "掘削工事"]

    def test_hierarchy_survives_a_summary_style_sheet(self, fake_pdf):
        """種目行にも「数量1・単価・金額」が入る様式で階層が潰れない。

        数量と単価がそろう行を細目別とする規則だけで判定していた頃は、
        集計行も明細行も区別できず全行が細目別になっていた。
        """
        fake_pdf(_ninomiya_pages())
        result = boq_import.parse_pdf(b"")

        assert [(d.name, d.level) for d in result.drafts] == [
            ("舗装工", BoqLine.Level.SHUMOKU),
            ("アスファルト撤去工", BoqLine.Level.KAMOKU),
            ("アスファルトカッター工", BoqLine.Level.CHUKAMOKU),
            ("施工費", BoqLine.Level.SAIMOKU),
            ("アスファルト撤去", BoqLine.Level.CHUKAMOKU),
            ("施工費", BoqLine.Level.SAIMOKU),
            ("掘削工事", BoqLine.Level.SHUMOKU),
            ("掘削工", BoqLine.Level.KAMOKU),
            # 3階層で終わる枝。末端で数量×単価を積むので細目別。
            ("施工費", BoqLine.Level.SAIMOKU),
        ]

    def test_hierarchy_continues_across_pages(self, fake_pdf):
        """ページを跨いだ続きの行を、枝の末端と取り違えない。

        ページごとに読むと、ページ末尾の行が「下に行が無い＝末端」に
        見えて細目別に落ちていた。
        """
        fake_pdf(_ninomiya_pages())
        result = boq_import.parse_pdf(b"")

        # p.2 の末尾（施工費）と p.3 の先頭（アスファルト撤去）
        assert result.drafts[3].level == BoqLine.Level.SAIMOKU
        assert result.drafts[4].level == BoqLine.Level.CHUKAMOKU

    def test_quantity_and_unit_are_read_from_one_cell(self, fake_pdf):
        fake_pdf(_ninomiya_pages())
        result = boq_import.parse_pdf(b"")

        meisai = [d for d in result.drafts if d.level == BoqLine.Level.SAIMOKU]
        assert [(d.quantity, d.unit) for d in meisai] == [
            (Decimal("10.4"), "m"),
            (Decimal("6.4"), "㎡"),
            (Decimal("2.88"), "㎥"),
        ]

    def test_amounts_match_quantity_times_unit_price(self, fake_pdf):
        fake_pdf(_ninomiya_pages())
        result = boq_import.parse_pdf(b"")

        for draft in result.drafts:
            if draft.level != BoqLine.MEISAI_LEVEL:
                continue
            assert round(draft.quantity * draft.unit_price) == draft.amount

    def test_cover_page_does_not_warn(self, fake_pdf):
        """表紙は内訳書のページではない。Excel の表紙シートと同じ扱い。"""
        pages = _ninomiya_pages()
        pages.insert(0, _FakePage("御見積書", []))
        fake_pdf(pages)

        assert boq_import.parse_pdf(b"").warnings == []


# ---------------------------------------------------------------------------
# 保存と階層の組み立て
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestLoadDrafts:
    def _drafts(self):
        return [
            boq_import.BoqDraft(level=BoqLine.Level.SHUMOKU, name="電気設備工事"),
            boq_import.BoqDraft(level=BoqLine.Level.KAMOKU, name="電灯設備"),
            boq_import.BoqDraft(
                level=BoqLine.Level.SAIMOKU, name="VVFケーブル",
                quantity=Decimal("1200"), unit_price=Decimal("150"),
            ),
            boq_import.BoqDraft(level=BoqLine.Level.KAMOKU, name="動力設備"),
            boq_import.BoqDraft(
                level=BoqLine.Level.SAIMOKU, name="CVケーブル",
                quantity=Decimal("300"), unit_price=Decimal("800"),
            ),
        ]

    def test_builds_parent_links(self, company_a):
        site = _site(company_a)
        created = boq_import.load_drafts(
            self._drafts(), company=company_a, site=site,
        )
        assert created == 5

        lines = {line.name: line for line in BoqLine.unscoped.filter(site=site)}
        assert lines["電気設備工事"].parent is None
        assert lines["電灯設備"].parent == lines["電気設備工事"]
        assert lines["VVFケーブル"].parent == lines["電灯設備"]
        # 2つめの科目の子が、1つめの科目の子にぶら下がらないこと。
        assert lines["動力設備"].parent == lines["電気設備工事"]
        assert lines["CVケーブル"].parent == lines["動力設備"]

    def test_amount_is_calculated(self, company_a):
        site = _site(company_a)
        boq_import.load_drafts(self._drafts(), company=company_a, site=site)
        line = BoqLine.unscoped.get(site=site, name="VVFケーブル")
        assert line.amount == Decimal("180000")

    def test_replace_clears_existing(self, company_a):
        site = _site(company_a)
        _line(company_a, site, "古い明細")
        boq_import.load_drafts(
            self._drafts(), company=company_a, site=site, replace=True,
        )
        names = set(BoqLine.unscoped.filter(site=site).values_list("name", flat=True))
        assert "古い明細" not in names

    def test_append_keeps_existing(self, company_a):
        site = _site(company_a)
        _line(company_a, site, "古い明細")
        boq_import.load_drafts(
            self._drafts(), company=company_a, site=site, replace=False,
        )
        names = set(BoqLine.unscoped.filter(site=site).values_list("name", flat=True))
        assert "古い明細" in names

    def test_requires_exactly_one_owner(self, company_a):
        with pytest.raises(ValueError):
            boq_import.load_drafts([], company=company_a)


@pytest.mark.django_db
class TestRebuildTree:
    def test_reassigns_parents_from_order(self, company_a):
        site = _site(company_a)
        parent = _line(company_a, site, "電灯設備", BoqLine.Level.KAMOKU, order=1)
        child = _line(company_a, site, "VVF", BoqLine.Level.SAIMOKU, order=2)
        # わざと誤った親子にしておく
        child.parent = None
        child.save()

        boq_import.rebuild_tree([parent, child])

        child.refresh_from_db()
        assert child.parent == parent


# ---------------------------------------------------------------------------
# 画面
# ---------------------------------------------------------------------------


@pytest.mark.django_db
class TestSiteDetailBoqSection:
    def test_section_is_rendered(self, client, user_a, company_a):
        site = _site(company_a)
        client.force_login(user_a)
        response = client.get(reverse("sites:detail", args=[site.pk]))

        assert response.status_code == 200
        assert "内訳書・内訳明細書" in response.content.decode()

    def test_lines_and_total_are_shown(self, client, user_a, company_a):
        site = _site(company_a)
        _line(
            company_a, site, "VVFケーブル", spec="1.6mm 2芯", unit="m",
            quantity=Decimal("1200"), unit_price=Decimal("150"),
            amount=Decimal("180000"),
        )
        client.force_login(user_a)
        response = client.get(reverse("sites:detail", args=[site.pk]))
        body = response.content.decode()

        assert "VVFケーブル" in body
        assert "1.6mm 2芯" in body
        assert response.context["boq_total"] == Decimal("180000")
        assert response.context["boq_meisai_count"] == 1

    def test_total_counts_top_level_only(self, client, user_a, company_a):
        """種目を立てた内訳書で、上位行と細目を二重に数えない。"""
        site = _site(company_a)
        parent = _line(
            company_a, site, "電気設備工事", BoqLine.Level.SHUMOKU,
            order=1, amount=Decimal("180000"),
        )
        _line(
            company_a, site, "VVFケーブル", BoqLine.Level.SAIMOKU,
            order=2, amount=Decimal("180000"), parent=parent,
        )
        client.force_login(user_a)
        response = client.get(reverse("sites:detail", args=[site.pk]))

        assert response.context["boq_total"] == Decimal("180000")


@pytest.mark.django_db
class TestSiteBoqEdit:
    def _post_data(self, rows, initial=0):
        data = {
            "form-TOTAL_FORMS": str(len(rows)),
            "form-INITIAL_FORMS": str(initial),
            "form-MIN_NUM_FORMS": "0",
            "form-MAX_NUM_FORMS": "1000",
        }
        for i, row in enumerate(rows):
            for key, value in row.items():
                data[f"form-{i}-{key}"] = value
        return data

    def test_creates_rows_from_table(self, client, user_a, company_a):
        site = _site(company_a)
        client.force_login(user_a)

        response = client.post(
            reverse("sites:boq_edit", args=[site.pk]),
            self._post_data([
                {"level": BoqLine.Level.KAMOKU, "name": "電灯設備"},
                {
                    "level": BoqLine.Level.SAIMOKU, "name": "VVFケーブル",
                    "spec": "1.6mm 2芯", "unit": "m",
                    "quantity": "1200", "unit_price": "150",
                },
                {"level": "", "name": ""},   # 空行
            ]),
        )

        assert response.status_code == 302
        lines = list(BoqLine.unscoped.filter(site=site).order_by("sort_order"))
        assert [line.name for line in lines] == ["電灯設備", "VVFケーブル"]
        assert lines[1].amount == Decimal("180000")
        assert lines[1].parent == lines[0]

    def test_blank_rows_do_not_error(self, client, user_a, company_a):
        site = _site(company_a)
        client.force_login(user_a)
        response = client.post(
            reverse("sites:boq_edit", args=[site.pk]),
            self._post_data([{"level": "", "name": ""} for _ in range(5)]),
        )
        assert response.status_code == 302
        assert BoqLine.unscoped.filter(site=site).count() == 0

    def test_level_defaults_from_numbers(self, client, user_a, company_a):
        """階層未指定でも、数量があれば内訳明細書として保存する。"""
        site = _site(company_a)
        client.force_login(user_a)
        client.post(
            reverse("sites:boq_edit", args=[site.pk]),
            self._post_data([
                {"level": "", "name": "VVFケーブル", "quantity": "10", "unit_price": "100"},
            ]),
        )
        line = BoqLine.unscoped.get(site=site)
        assert line.level == BoqLine.MEISAI_LEVEL

    def test_saved_line_belongs_to_site_not_project(self, client, user_a, company_a):
        site = _site(company_a)
        client.force_login(user_a)
        client.post(
            reverse("sites:boq_edit", args=[site.pk]),
            self._post_data([{"level": BoqLine.Level.SAIMOKU, "name": "VVF"}]),
        )
        line = BoqLine.unscoped.get(site=site)
        assert line.project_id is None
        assert line.company_id == company_a.id


@pytest.mark.django_db
class TestSiteBoqImportView:
    def test_uploads_excel_and_redirects(self, client, user_a, company_a):
        from django.core.files.uploadedfile import SimpleUploadedFile

        site = _site(company_a)
        client.force_login(user_a)
        data = _excel_bytes([
            HEADER,
            ["細目別", "VVFケーブル", "1.6-2C", "m", 1200, 150, 180000, ""],
        ])
        upload = SimpleUploadedFile(
            "内訳書.xlsx", data,
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )

        response = client.post(
            reverse("sites:boq_import", args=[site.pk]),
            {"upload": upload, "replace": "on"},
        )

        assert response.status_code == 302
        line = BoqLine.unscoped.get(site=site)
        assert line.name == "VVFケーブル"
        assert line.amount == Decimal("180000")

    def test_unreadable_file_stays_on_page(self, client, user_a, company_a):
        from django.core.files.uploadedfile import SimpleUploadedFile

        site = _site(company_a)
        client.force_login(user_a)
        upload = SimpleUploadedFile("内訳書.xlsx", b"broken")

        response = client.post(
            reverse("sites:boq_import", args=[site.pk]),
            {"upload": upload},
        )

        assert response.status_code == 200
        assert BoqLine.unscoped.filter(site=site).count() == 0


@pytest.mark.django_db
class TestSiteBoqExport:
    def test_downloads_excel(self, client, user_a, company_a):
        site = _site(company_a)
        _line(
            company_a, site, "VVFケーブル", unit="m",
            quantity=Decimal("1200"), unit_price=Decimal("150"),
            amount=Decimal("180000"),
        )
        client.force_login(user_a)
        response = client.get(reverse("sites:boq_export", args=[site.pk]))

        assert response.status_code == 200
        assert "spreadsheetml" in response["Content-Type"]

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        # 名称セルは階層ぶんのインデントが付くので、前後の空白を落として比べる。
        values = [
            str(cell.value).strip()
            for row in wb["内訳書"].iter_rows()
            for cell in row
            if cell.value is not None
        ]
        assert "VVFケーブル" in values
        assert site.name in values
        assert "180000" in values or 180000 in [
            cell.value
            for row in wb["内訳書"].iter_rows()
            for cell in row
        ]

    def test_export_requires_one_owner(self, company_a):
        with pytest.raises(ValueError):
            export_boq_to_excel()

    def test_roundtrip_export_then_import(self, company_a):
        """出力した Excel を読み戻せる（階層列を自分で書いているため）。"""
        site = _site(company_a)
        parent = _line(
            company_a, site, "電灯設備", BoqLine.Level.KAMOKU,
            order=1, amount=Decimal("180000"),
        )
        _line(
            company_a, site, "VVFケーブル", BoqLine.Level.SAIMOKU, order=2,
            spec="1.6-2C", unit="m", quantity=Decimal("1200"),
            unit_price=Decimal("150"), amount=Decimal("180000"), parent=parent,
        )

        result = boq_import.parse_excel(export_boq_to_excel(site=site))

        names = [d.name for d in result.drafts]
        assert "電灯設備" in names
        assert "VVFケーブル" in names
        meisai = next(d for d in result.drafts if d.name == "VVFケーブル")
        assert meisai.level == BoqLine.Level.SAIMOKU
        assert meisai.quantity == Decimal("1200")
