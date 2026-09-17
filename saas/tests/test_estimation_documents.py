"""積算案件の入力補助と、案件資料・日程の読み取り（ADR-0080）。

確かめること:

- 発注機関は候補に無くても手で入力でき、そのとき発注機関マスタにも登録される
- 積算担当者は手入力で、過去に入れた名前が候補に出る
- 工期末日を西暦・日付まで出す
- 案件資料（PDF・Excel）を登録・閲覧・削除できる
- 資料の PDF から入札までの日程を読み取り、工程として取り込める
- 入札から引っ越した案件は、公告のURLに詳細から飛べる
- 以上が他社のデータに触れない（テナント越境）
"""
import datetime
import io

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.bids.models import BidProject
from apps.estimation.models import (
    EstimationDocument,
    EstimationPhase,
    EstimationProject,
    Orderer,
)
from apps.estimation.services.document_schedule import (
    ScheduleReadError,
    import_as_phases,
    read_schedule,
    schedule_rows,
)
from apps.estimation.services.from_bid import resolve_or_create_orderer

# 公告の本文。本物は pdfplumber で読むが、テストでは読み取り後の文字列から先を確かめる。
# 200文字（announcement.MIN_TEXT_CHARS）を超える長さにしておく。
# これを下回ると「スキャン画像で文字が無い」と見なされ、AI での読み取りに回る。
ANNOUNCEMENT_TEXT = """
入札公告

次のとおり一般競争入札に付します。

１．工事概要
(1) 工事名 七沢自然ふれあいセンター宿泊棟長寿命化（機能回復・冷暖房設備）改修（電気）工事
(2) 工事場所 厚木市七沢2440番地
(3) 工事内容 宿泊棟の電気設備一式の改修。受変電設備、動力設備、電灯設備、
    自動火災報知設備の更新を行う。
(4) 工期 契約の日から令和9年6月30日まで

２．競争参加資格
(1) 神奈川県の入札参加資格において電気工事の認定を受けていること。
(2) 建設業法による電気工事業の許可を受けていること。

４．入札手続等
① 申請書及び資料の提出期限 令和8年9月17日 正午
② 入札書の受領期限 令和8年10月27日 17時
③ 開札の日時及び場所 令和8年11月11日 15時 第2会議室
"""


def _pdf(name="公告.pdf"):
    body = b"%PDF-1.7\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"
    return SimpleUploadedFile(name, body, content_type="application/pdf")


def _xlsx(name="内訳書.xlsx"):
    buffer = io.BytesIO()
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "数量"
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


def _orderer(company, name="厚木市"):
    orderer, _ = Orderer.unscoped.get_or_create(company=company, name=name)
    return orderer


def _project(company, name="七沢センター改修", **kwargs):
    fields = {"orderer": _orderer(company), "status": EstimationProject.Status.ESTIMATING}
    fields.update(kwargs)
    return EstimationProject.unscoped.create(company=company, name=name, **fields)


def _document(project, **kwargs):
    fields = {
        "name": "公告",
        "doc_type": EstimationDocument.DocType.ANNOUNCEMENT,
        "kind": EstimationDocument.Kind.PDF,
        "original_filename": "公告.pdf",
    }
    fields.update(kwargs)
    document = EstimationDocument(company=project.company, project=project, **fields)
    document.file = _pdf()
    document.save()
    return document


def _form_post(project, **overrides):
    """積算案件の編集画面に送る中身。"""
    data = {
        "name": project.name,
        "orderer_name": project.orderer.name,
        "standard": "", "site": "", "bid_project": "",
        "primary_work_category": project.primary_work_category,
        "status": project.status,
        "bid_announcement_date": "", "bid_opening_date": "",
        "construction_period_days": "", "construction_end_date": "",
        "estimator_name": project.estimator_name,
        "bid_amount": "", "award_amount": "", "notes": "",
    }
    data.update(overrides)
    return data


@pytest.mark.django_db
class TestOrdererFreeText:
    def test_候補に無い機関名を打つと発注機関マスタにも登録される(
        self, client, company_a, user_a,
    ):
        project = _project(company_a)
        client.force_login(user_a)

        res = client.post(
            f"/estimation/projects/{project.pk}/edit/",
            _form_post(project, orderer_name="神奈川県企業局"),
        )

        project.refresh_from_db()
        assert res.status_code == 302
        assert project.orderer.name == "神奈川県企業局"
        assert Orderer.unscoped.filter(
            company=company_a, name="神奈川県企業局",
        ).exists()

    def test_同じ機関名なら既存のマスタに寄る(self, company_a, user_a):
        existing = _orderer(company_a, "厚木市")

        found = resolve_or_create_orderer(company_a, "厚木市", created_by=user_a)

        assert found.pk == existing.pk
        assert Orderer.unscoped.filter(company=company_a, name="厚木市").count() == 1

    def test_前後の空白は落として引き当てる(self, company_a):
        existing = _orderer(company_a, "厚木市")

        assert resolve_or_create_orderer(company_a, "  厚木市  ").pk == existing.pk

    def test_空欄では作らない(self, company_a):
        assert resolve_or_create_orderer(company_a, "") is None
        assert resolve_or_create_orderer(company_a, "   ") is None

    def test_似た名前でも別の機関として残す(self, company_a):
        """「厚木市」と「厚木市教育委員会」は積算基準が違うことがある。

        顧客マスタのような表記ゆれの正規化をすると取り違える。
        """
        resolve_or_create_orderer(company_a, "厚木市")
        resolve_or_create_orderer(company_a, "厚木市教育委員会")

        assert Orderer.unscoped.filter(company=company_a).count() == 2

    def test_編集画面に候補が並ぶ(self, client, company_a, user_a):
        _orderer(company_a, "厚木市")
        project = _project(company_a)
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/edit/").content.decode()

        assert 'id="orderer-name-options"' in html
        assert '<option value="厚木市">' in html
        assert 'list="orderer-name-options"' in html


@pytest.mark.django_db
class TestEstimatorAndEndDate:
    def test_担当者を手で入れられる(self, client, company_a, user_a):
        project = _project(company_a)
        client.force_login(user_a)

        client.post(
            f"/estimation/projects/{project.pk}/edit/",
            _form_post(project, estimator_name="佐野 歩"),
        )

        project.refresh_from_db()
        assert project.estimator_name == "佐野 歩"

    def test_過去に入れた担当者が候補に出る(self, client, company_a, user_a):
        _project(company_a, name="前の案件", estimator_name="佐野 歩")
        _project(company_a, name="別の案件", estimator_name="片桐 徳男")
        project = _project(company_a)
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/edit/").content.decode()

        options = html.split('id="estimator-name-options"')[1].split("</datalist>")[0]
        assert "佐野 歩" in options
        assert "片桐 徳男" in options

    def test_同じ担当者は候補に1つだけ出る(self, client, company_a, user_a):
        _project(company_a, name="案件1", estimator_name="佐野 歩")
        _project(company_a, name="案件2", estimator_name="佐野 歩")
        project = _project(company_a)
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/edit/").content.decode()

        options = html.split('id="estimator-name-options"')[1].split("</datalist>")[0]
        assert options.count("佐野 歩") == 1

    def test_他社の担当者は候補に出ない(self, client, company_a, company_b, user_a):
        _project(company_b, name="B社の案件", estimator_name="他社の人")
        project = _project(company_a)
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/edit/").content.decode()

        assert "他社の人" not in html

    def test_工期末日を西暦で日付まで出す(self, client, company_a, user_a):
        project = _project(
            company_a, construction_end_date=datetime.date(2027, 6, 30),
            construction_period_days=404,
        )
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/").content.decode()

        assert "2027年6月30日 まで" in html
        assert "404日" in html

    def test_工期末日が開札予定日より前なら保存しない(self, client, company_a, user_a):
        project = _project(company_a)
        client.force_login(user_a)

        res = client.post(
            f"/estimation/projects/{project.pk}/edit/",
            _form_post(
                project,
                bid_opening_date="2026-11-11",
                construction_end_date="2026-10-01",
            ),
        )

        project.refresh_from_db()
        assert res.status_code == 200
        assert project.construction_end_date is None


@pytest.mark.django_db
class TestDocuments:
    def test_資料を登録できる(self, client, company_a, user_a):
        project = _project(company_a)
        client.force_login(user_a)

        res = client.post(f"/estimation/projects/{project.pk}/documents/add/", {
            "name": "入札公告", "doc_type": "announcement",
            "file": _pdf(), "memo": "厚木市のサイトから",
        })

        assert res.status_code == 302
        document = EstimationDocument.unscoped.get(company=company_a, project=project)
        assert document.name == "入札公告"
        assert document.kind == EstimationDocument.Kind.PDF
        assert document.original_filename == "公告.pdf"
        assert document.size > 0

    def test_資料名が空ならファイル名を使う(self, client, company_a, user_a):
        project = _project(company_a)
        client.force_login(user_a)

        client.post(f"/estimation/projects/{project.pk}/documents/add/", {
            "name": "", "doc_type": "spec", "file": _pdf("仕様書.pdf"),
        })

        document = EstimationDocument.unscoped.get(company=company_a)
        assert document.name == "仕様書"

    def test_Excelも登録できる(self, client, company_a, user_a):
        project = _project(company_a)
        client.force_login(user_a)

        client.post(f"/estimation/projects/{project.pk}/documents/add/", {
            "name": "内訳書", "doc_type": "boq", "file": _xlsx(),
        })

        document = EstimationDocument.unscoped.get(company=company_a)
        assert document.kind == EstimationDocument.Kind.EXCEL
        assert document.is_pdf is False

    def test_拡張子だけ変えたファイルは弾く(self, client, company_a, user_a):
        project = _project(company_a)
        client.force_login(user_a)

        client.post(f"/estimation/projects/{project.pk}/documents/add/", {
            "name": "偽物", "doc_type": "other",
            "file": SimpleUploadedFile("偽物.pdf", b"not a pdf"),
        })

        assert not EstimationDocument.unscoped.filter(company=company_a).exists()

    def test_資料を開ける(self, client, company_a, user_a):
        document = _document(_project(company_a))
        client.force_login(user_a)

        res = client.get(f"/estimation/documents/{document.pk}/file/")

        assert res.status_code == 200
        assert res["Cache-Control"] == "private, no-cache"

    def test_資料を消すとファイルも消える(
        self, client, company_a, user_a, django_capture_on_commit_callbacks,
    ):
        """ファイルの削除は transaction.on_commit に入れてある。

        テストはトランザクションの中で走るのでそのままでは発火しない。
        自社書類のテストと同じく、コミット後の処理を明示的に走らせる。
        """
        document = _document(_project(company_a))
        storage, name = document.file.storage, document.file.name
        client.force_login(user_a)

        with django_capture_on_commit_callbacks(execute=True):
            client.post(f"/estimation/documents/{document.pk}/delete/")

        assert not EstimationDocument.unscoped.filter(pk=document.pk).exists()
        assert not storage.exists(name)

    def test_詳細画面に資料が並ぶ(self, client, company_a, user_a):
        project = _project(company_a)
        _document(project, name="入札公告")
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/").content.decode()

        assert "案件資料" in html
        assert "入札公告" in html


@pytest.mark.django_db
class TestScheduleFromDocument:
    def test_PDFから日程を読み取る(self, monkeypatch, company_a, user_a):
        document = _document(_project(company_a))
        monkeypatch.setattr(
            "apps.bids.announcement.extract_text", lambda data: ANNOUNCEMENT_TEXT,
        )

        result = read_schedule(document, user=user_a)

        document.refresh_from_db()
        assert result["source"] == "text"
        assert len(result["schedule"]) >= 3
        assert document.ai_checked_at is not None
        labels = [row["label"] for row in schedule_rows(document)]
        assert any("入札書" in label for label in labels)
        assert any("開札" in label for label in labels)

    def test_読み取った日程を工程に取り込む(self, monkeypatch, company_a, user_a):
        project = _project(company_a)
        document = _document(project)
        monkeypatch.setattr(
            "apps.bids.announcement.extract_text", lambda data: ANNOUNCEMENT_TEXT,
        )
        read_schedule(document, user=user_a)

        count = import_as_phases(document, created_by=user_a)

        phases = list(EstimationPhase.unscoped.filter(project=project))
        assert count == len(phases) >= 3
        # 期限の早い順に並び、どの工程も開始日と終了日が入る
        for phase in phases:
            assert phase.start_date and phase.end_date
            assert phase.start_date <= phase.end_date
        opening = next(p for p in phases if "開札" in p.name)
        assert opening.end_date == datetime.date(2026, 11, 11)

    def test_取り込み直しても増えない(self, monkeypatch, company_a, user_a):
        project = _project(company_a)
        document = _document(project)
        monkeypatch.setattr(
            "apps.bids.announcement.extract_text", lambda data: ANNOUNCEMENT_TEXT,
        )
        read_schedule(document, user=user_a)

        first = import_as_phases(document, created_by=user_a)
        second = import_as_phases(document, created_by=user_a)

        assert first >= 3
        assert second == 0
        assert EstimationPhase.unscoped.filter(project=project).count() == first

    def test_取り込み直しても手で直した日付は戻らない(self, monkeypatch, company_a, user_a):
        project = _project(company_a)
        document = _document(project)
        monkeypatch.setattr(
            "apps.bids.announcement.extract_text", lambda data: ANNOUNCEMENT_TEXT,
        )
        read_schedule(document, user=user_a)
        import_as_phases(document, created_by=user_a)
        phase = EstimationPhase.unscoped.filter(project=project).first()
        phase.end_date = datetime.date(2026, 12, 25)
        phase.save()

        import_as_phases(document, created_by=user_a)

        phase.refresh_from_db()
        assert phase.end_date == datetime.date(2026, 12, 25)

    def test_Excelは読み取れないと伝える(self, company_a):
        project = _project(company_a)
        document = EstimationDocument(
            company=company_a, project=project, name="内訳書",
            doc_type=EstimationDocument.DocType.BOQ,
            kind=EstimationDocument.Kind.EXCEL, original_filename="内訳書.xlsx",
        )
        document.file = _xlsx()
        document.save()

        with pytest.raises(ScheduleReadError, match="PDF"):
            read_schedule(document)

    def test_日程が無いPDFは読み取れないと伝える(self, monkeypatch, company_a):
        document = _document(_project(company_a))
        monkeypatch.setattr(
            "apps.bids.announcement.extract_text", lambda data: "日程の書かれていない書類",
        )
        monkeypatch.setattr("apps.bids.announcement_llm.is_available", lambda: False)

        with pytest.raises(ScheduleReadError, match="読み取れませんでした"):
            read_schedule(document)

    def test_画面から読み取って取り込める(self, monkeypatch, client, company_a, user_a):
        project = _project(company_a)
        document = _document(project)
        monkeypatch.setattr(
            "apps.bids.announcement.extract_text", lambda data: ANNOUNCEMENT_TEXT,
        )
        client.force_login(user_a)

        client.post(f"/estimation/documents/{document.pk}/read-schedule/")
        html = client.get(f"/estimation/projects/{project.pk}/").content.decode()
        assert "読み取った入札までの日程" in html

        client.post(f"/estimation/documents/{document.pk}/import-schedule/")
        assert EstimationPhase.unscoped.filter(project=project).exists()


@pytest.mark.django_db
class TestBidAnnouncementLink:
    def test_公告のURLに詳細から飛べる(self, client, company_a, user_a):
        bid = BidProject.unscoped.create(
            company=company_a, title="七沢センター改修",
            source_url="https://example.jp/koukoku.pdf",
            document_urls="https://example.jp/koukoku.pdf\nhttps://example.jp/betsu.pdf",
        )
        project = _project(company_a, bid_project=bid)
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/").content.decode()

        assert "公告PDFを開く" in html
        # source_url と同じものは「公開文書」に二重で出さない
        assert html.count("betsu.pdf") >= 1
        assert "公開文書 1" in html

    def test_公告のURLが無いときは案内を出す(self, client, company_a, user_a):
        bid = BidProject.unscoped.create(company=company_a, title="URLの無い案件")
        project = _project(company_a, bid_project=bid)
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/").content.decode()

        assert "公告のURLが入っていません" in html


@pytest.mark.django_db
class TestSingleGantt:
    def test_詳細にこの案件の流れが出る(self, company_a, client, user_a):
        project = _project(company_a)
        EstimationPhase.unscoped.create(
            company=company_a, project=project, name="参加申請",
            start_date=datetime.date(2026, 9, 1),
            end_date=datetime.date(2026, 9, 17),
        )
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/").content.decode()

        assert "この案件の流れ" in html
        assert "est-gantt-seg" in html

    def test_他の案件は混ざらない(self, company_a, client, user_a):
        project = _project(company_a)
        EstimationPhase.unscoped.create(
            company=company_a, project=project, name="参加申請",
            start_date=datetime.date(2026, 9, 1),
            end_date=datetime.date(2026, 9, 17),
        )
        other = _project(company_a, name="別の案件")
        EstimationPhase.unscoped.create(
            company=company_a, project=other, name="参加申請",
            start_date=datetime.date(2026, 9, 1),
            end_date=datetime.date(2026, 9, 17),
        )
        client.force_login(user_a)

        html = client.get(f"/estimation/projects/{project.pk}/").content.decode()

        flow = html.split("この案件の流れ")[1].split("案件資料")[0]
        assert "別の案件" not in flow


@pytest.mark.django_db
class TestTenantIsolation:
    def test_他社の案件に資料を足せない(self, client, company_a, company_b, user_a):
        other = _project(company_b, name="B社の案件")
        client.force_login(user_a)

        res = client.post(f"/estimation/projects/{other.pk}/documents/add/", {
            "name": "割り込み", "doc_type": "other", "file": _pdf(),
        })

        assert res.status_code == 404
        assert not EstimationDocument.unscoped.filter(name="割り込み").exists()

    def test_他社の資料は開けない(self, client, company_a, company_b, user_a):
        document = _document(_project(company_b))
        client.force_login(user_a)

        assert client.get(f"/estimation/documents/{document.pk}/file/").status_code == 404

    def test_他社の資料は消せない(self, client, company_a, company_b, user_a):
        document = _document(_project(company_b))
        client.force_login(user_a)

        res = client.post(f"/estimation/documents/{document.pk}/delete/")

        assert res.status_code == 404
        assert EstimationDocument.unscoped.filter(pk=document.pk).exists()

    def test_他社の資料を読み取れない(self, client, company_a, company_b, user_a):
        document = _document(_project(company_b))
        client.force_login(user_a)

        res = client.post(f"/estimation/documents/{document.pk}/read-schedule/")

        assert res.status_code == 404

    def test_ログインしていないと資料を開けない(self, client, company_a):
        document = _document(_project(company_a))

        res = client.get(f"/estimation/documents/{document.pk}/file/")

        assert res.status_code == 302
