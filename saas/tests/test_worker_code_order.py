"""作業員一覧・社員名簿の社員番号順ソートのテスト。"""

from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.core.tenant_context import set_current_company
from apps.workers.models import Worker, employee_code_sort_key


def test_sort_key_groups_by_alphabet_then_number():
    codes = ["E10", "F1", "E2", "e1", "A100", "E010"]
    assert sorted(codes, key=employee_code_sort_key) == ["e1", "E2", "E010", "E10", "A100", "F1"]


def test_sort_key_uses_fixed_prefix_order():
    codes = ["G1", "A1", "P1", "T1", "E1", "S1", "Y1", "B1", "F1", "Z1"]
    assert sorted(codes, key=employee_code_sort_key) == [
        "Y1", "S1", "E1", "T1", "P1", "A1", "G1", "B1", "F1", "Z1",
    ]


def test_sort_key_handles_separators_and_missing_number():
    codes = ["E-3", "E_1", "E 2", "E"]
    assert sorted(codes, key=employee_code_sort_key) == ["E", "E_1", "E 2", "E-3"]


def test_sort_key_puts_blank_last():
    codes = ["", None, "B1", "A1"]
    assert sorted(codes, key=employee_code_sort_key) == ["A1", "B1", "", None]


@pytest.fixture
def workers(company_a):
    set_current_company(company_a)
    rows = [
        ("E10", "十郎", "ジュウロウ", True),
        ("A1", "太郎", "タロウ", True),
        ("E2", "次郎", "ジロウ", True),
        ("", "花子", "ハナコ", True),
        ("S5", "五郎", "ゴロウ", False),
        ("", "梅子", "ウメコ", True),
        ("Y3", "三郎", "サブロウ", True),
    ]
    created = [
        Worker.objects.create(
            company=company_a, employee_code=code, name=name, name_kana=kana,
            is_active=active, hourly_cost=3000,
        )
        for code, name, kana, active in rows
    ]
    yield created
    set_current_company(None)


@pytest.mark.django_db
def test_list_orders_by_employee_code(client, user_a, workers):
    client.force_login(user_a)
    res = client.get(reverse("workers:list") + "?inactive=1")
    assert res.status_code == 200
    active = [w.employee_code or w.name for w in res.context["active_workers"]]
    inactive = [w.employee_code for w in res.context["inactive_workers"]]
    assert active == ["Y3", "E2", "E10", "A1", "梅子", "花子"]
    assert inactive == ["S5"]


@pytest.mark.django_db
def test_excel_orders_by_employee_code(client, user_a, workers):
    client.force_login(user_a)
    res = client.get(reverse("workers:excel") + "?active_only=0")
    assert res.status_code == 200
    ws = load_workbook(BytesIO(res.content)).active
    codes = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert codes == ["Y3", "S5", "E2", "E10", "A1", None, None]
