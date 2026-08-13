"""見積ファイル（ライデンの CSV / 見積書の PDF）からの現場取り込み。

読み取りは「ラベル名で探す」方式なので、列位置ではなくラベルの表記ゆれと
値の置き場所（同セル / 右 / 下）を固定する。読めなかった項目は推測せず
空で返し、確認画面で人が入力できることも合わせて確かめる。
"""

import datetime

import pytest
from django.urls import reverse

from apps.core.tenant_context import set_current_company
from apps.masters.models import Customer
from apps.sites import importer
from apps.sites.importer import (
    clean_company_name,
    normalize_company_name,
    parse_estimate_file,
)
from apps.sites.models import EstimateImport, Site
from apps.sites.services import find_customer_by_name

# 見出し部だけを持つ、ライデン出力を模した CSV。
ESTIMATE_CSV = """御見積書,,,,
株式会社サンプル建設 御中,,,,
,,見積番号,Q-2026-0142,
,,見積日,2026/08/12,
工事件名,○○ビル 電気設備改修工事,,,
施工場所,山形県酒田市中町1-2-3,,,
工期,2026/09/01～2026/11/30,,,
支払条件,月末締め翌月末現金払い,,,
見積有効期限,2026/09/30,,,
備考,足場は別途,,,
見積金額,"3,480,000円",,,
担当者,剣持,,,
工事区分,電気設備工事,,,
,,,,
No,名称,数量,単位,金額
1,電線管,120,m,240000
"""


def _write_csv(tmp_path, text, encoding="cp932", name="estimate.csv"):
    path = tmp_path / name
    path.write_text(text, encoding=encoding)
    return path


# ---------------------------------------------------------------------------
# CSV の読み取り
# ---------------------------------------------------------------------------

class TestParseCsv:
    def test_reads_every_header_field(self, tmp_path):
        data = parse_estimate_file(_write_csv(tmp_path, ESTIMATE_CSV), ".csv")

        assert data["code"] == "Q-2026-0142"
        assert data["name"] == "○○ビル 電気設備改修工事"
        assert data["customer_name"] == "株式会社サンプル建設"
        assert data["payment_terms"] == "月末締め翌月末現金払い"
        assert data["address"] == "山形県酒田市中町1-2-3"
        assert data["contract_amount"] == 3480000
        assert data["start_date"] == "2026-09-01"
        assert data["end_date"] == "2026-11-30"
        assert data["estimate_valid_until"] == "2026/09/30"
        assert data["note"] == "足場は別途"
        assert data["missing"] == []

    def test_collects_labels_it_was_not_told_about(self, tmp_path):
        # 「担当者」「工事区分」は決め打ちのラベルに無い。取りこぼさず拾うこと。
        data = parse_estimate_file(_write_csv(tmp_path, ESTIMATE_CSV), ".csv")
        details = dict(data["details"])

        assert details["担当者"] == "剣持"
        assert details["工事区分"] == "電気設備工事"

    def test_detail_rows_are_not_collected_as_pairs(self, tmp_path):
        # 明細行（非空セルが多い行）まで拾うとノイズに埋もれる。
        data = parse_estimate_file(_write_csv(tmp_path, ESTIMATE_CSV), ".csv")
        labels = [label for label, _ in data["details"]]

        assert "電線管" not in labels
        assert "名称" not in labels

    def test_values_already_stored_in_fields_are_not_repeated(self, tmp_path):
        data = parse_estimate_file(_write_csv(tmp_path, ESTIMATE_CSV), ".csv")
        values = [value for _, value in data["details"]]

        assert "Q-2026-0142" not in values
        assert "月末締め翌月末現金払い" not in values

    def test_onchu_row_is_not_collected_as_a_pair(self, tmp_path):
        csv_text = "株式会社北日本建設,御中\n担当,山田\n"
        data = parse_estimate_file(_write_csv(tmp_path, csv_text), ".csv")

        assert ("株式会社北日本建設", "御中") not in data["details"]
        assert ("担当", "山田") in data["details"]

    def test_reads_utf8_as_well_as_cp932(self, tmp_path):
        # ライデンの書き出しは CP932 が多いが、版によって UTF-8 のこともある。
        path = _write_csv(tmp_path, ESTIMATE_CSV, encoding="utf-8-sig", name="u.csv")
        assert parse_estimate_file(path, ".csv")["code"] == "Q-2026-0142"

    def test_value_written_in_the_same_cell(self, tmp_path):
        csv_text = "見積番号：Q-9999\nお支払い条件: 翌々月10日 手形120日\n"
        data = parse_estimate_file(_write_csv(tmp_path, csv_text), ".csv")

        assert data["code"] == "Q-9999"
        assert data["payment_terms"] == "翌々月10日 手形120日"

    def test_value_written_in_the_row_below(self, tmp_path):
        csv_text = "見積番号,件名\nQ-1234,倉庫新築電気工事\n"
        data = parse_estimate_file(_write_csv(tmp_path, csv_text), ".csv")

        assert data["code"] == "Q-1234"
        assert data["name"] == "倉庫新築電気工事"

    def test_fullwidth_and_alternate_labels(self, tmp_path):
        # 「見積Ｎｏ．」「御支払条件」のような表記でも拾えること。
        csv_text = "見積Ｎｏ．,Ｑ－５５５\n御支払条件,現金\n工事名称,テスト工事\n"
        data = parse_estimate_file(_write_csv(tmp_path, csv_text), ".csv")

        assert data["code"] == "Q-555"
        assert data["payment_terms"] == "現金"
        assert data["name"] == "テスト工事"

    def test_unreadable_fields_are_left_empty_not_guessed(self, tmp_path):
        csv_text = "件名,配線工事\n"
        data = parse_estimate_file(_write_csv(tmp_path, csv_text), ".csv")

        assert data["name"] == "配線工事"
        assert data["code"] is None
        assert data["payment_terms"] is None
        assert data["customer_name"] is None
        assert "見積番号（現場コード）" not in data["found"]
        assert "code" in data["missing"]
        assert "payment_terms" in data["missing"]

    def test_overlong_value_is_discarded_rather_than_stored(self, tmp_path):
        # 明細行などを掴んでしまった場合は、間違った値を入れるより空にする。
        csv_text = "支払条件," + "あ" * 300 + "\n"
        data = parse_estimate_file(_write_csv(tmp_path, csv_text), ".csv")

        assert data["payment_terms"] is None

    def test_rejects_unsupported_extension(self, tmp_path):
        path = tmp_path / "estimate.docx"
        path.write_text("dummy", encoding="utf-8")
        with pytest.raises(ValueError, match="CSV"):
            parse_estimate_file(path, ".docx")

    def test_old_excel_format_explains_how_to_convert(self, tmp_path):
        # openpyxl は .xls を読めない。黙って失敗させず直し方を伝える。
        path = tmp_path / "estimate.xls"
        path.write_bytes(b"dummy")
        with pytest.raises(ValueError, match=r"\.xlsx で保存し直して"):
            parse_estimate_file(path, ".xls")


# ---------------------------------------------------------------------------
# Excel の読み取り
# ---------------------------------------------------------------------------

def _write_xlsx(tmp_path, sheets, name="estimate.xlsx"):
    """{シート名: 行のリスト} から xlsx を作る。"""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        for row in rows:
            sheet.append(row)
    path = tmp_path / name
    workbook.save(path)
    return path


ESTIMATE_SHEET = [
    ["御見積書"],
    ["株式会社サンプル建設 御中"],
    [None, None, "見積番号", "Q-2026-0142"],
    ["工事件名", "○○ビル 電気設備改修工事"],
    ["施工場所", "山形県酒田市中町1-2-3"],
    ["工期", "2026/09/01～2026/11/30"],
    ["支払条件", "月末締め翌月末現金払い"],
    ["見積有効期限", "発行後30日間"],
    ["備考", "足場は別途"],
    ["見積金額", 3480000],
]


class TestParseExcel:
    def test_reads_every_header_field(self, tmp_path):
        path = _write_xlsx(tmp_path, {"表紙": ESTIMATE_SHEET})
        data = parse_estimate_file(path, ".xlsx")

        assert data["code"] == "Q-2026-0142"
        assert data["name"] == "○○ビル 電気設備改修工事"
        assert data["customer_name"] == "株式会社サンプル建設"
        assert data["payment_terms"] == "月末締め翌月末現金払い"
        assert data["address"] == "山形県酒田市中町1-2-3"
        assert data["contract_amount"] == 3480000
        assert data["start_date"] == "2026-09-01"
        assert data["end_date"] == "2026-11-30"
        # 日付とは限らないので、書かれたまま残す。
        assert data["estimate_valid_until"] == "発行後30日間"
        assert data["note"] == "足場は別途"
        assert data["missing"] == []

    def test_collects_unknown_labels_from_every_sheet(self, tmp_path):
        path = _write_xlsx(tmp_path, {
            "表紙": ESTIMATE_SHEET,
            "条件": [["工事区分", "電気設備工事"], ["現場代理人", "剣持"]],
        })
        details = dict(parse_estimate_file(path, ".xlsx")["details"])

        assert details["工事区分"] == "電気設備工事"
        assert details["現場代理人"] == "剣持"

    def test_reads_the_header_sheet_even_when_details_follow(self, tmp_path):
        # 表紙に見出し、別シートに明細、という作りでも読めること。
        path = _write_xlsx(tmp_path, {
            "表紙": ESTIMATE_SHEET,
            "内訳": [["No", "名称", "数量", "単位", "金額"], [1, "電線管", 120, "m", 240000]],
        })
        data = parse_estimate_file(path, ".xlsx")

        assert data["code"] == "Q-2026-0142"
        assert data["payment_terms"] == "月末締め翌月末現金払い"

    def test_date_cell_is_read_as_a_date(self, tmp_path):
        # 工期が文字列ではなく日付セルで入っている場合。
        path = _write_xlsx(tmp_path, {
            "表紙": [["工期", datetime.datetime(2026, 9, 1)]],
        })
        data = parse_estimate_file(path, ".xlsx")

        assert data["start_date"] == "2026-09-01"
        assert data["end_date"] is None

    def test_uncalculated_formula_is_left_empty(self, tmp_path):
        # 数式のまま保存されたファイルは data_only では値が取れない。
        # 誤った金額を入れるより空にして手入力してもらう。
        import openpyxl

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["見積金額", "=SUM(B2:B9)"])
        path = tmp_path / "formula.xlsx"
        workbook.save(path)

        assert parse_estimate_file(path, ".xlsx")["contract_amount"] is None

    def test_xlsm_is_accepted(self, tmp_path):
        path = _write_xlsx(tmp_path, {"表紙": ESTIMATE_SHEET}, name="estimate.xlsm")
        assert parse_estimate_file(path, ".xlsm")["code"] == "Q-2026-0142"


# ---------------------------------------------------------------------------
# 「御中」からの取引先名
# ---------------------------------------------------------------------------

class TestCustomerName:
    def test_company_and_onchu_in_one_cell(self, tmp_path):
        csv_text = "有限会社山田電機 御中\n"
        data = parse_estimate_file(_write_csv(tmp_path, csv_text), ".csv")
        assert data["customer_name"] == "有限会社山田電機"

    def test_company_and_onchu_in_separate_cells(self, tmp_path):
        csv_text = "株式会社北日本建設,御中\n"
        data = parse_estimate_file(_write_csv(tmp_path, csv_text), ".csv")
        assert data["customer_name"] == "株式会社北日本建設"

    def test_strips_honorifics_and_brackets(self):
        assert clean_company_name("　株式会社テスト　御中　") == "株式会社テスト"
        assert clean_company_name("テスト工務店 様") == "テスト工務店"

    @pytest.mark.parametrize("written", [
        "株式会社ABC", "(株)ABC", "㈱ABC", "㈱ ABC", "ＡＢＣ株式会社", "ABC(株)",
    ])
    def test_corporate_prefix_variants_normalize_to_the_same_key(self, written):
        assert normalize_company_name(written) == normalize_company_name("ABC")


# ---------------------------------------------------------------------------
# 得意先マスタへの引き当て
# ---------------------------------------------------------------------------

class TestFindCustomer:
    def test_matches_across_corporate_prefix_notation(self, company_a):
        customer = Customer.unscoped.create(
            company=company_a, code="C001", name="株式会社サンプル建設",
        )
        assert find_customer_by_name(company_a, "(株)サンプル建設") == customer
        assert find_customer_by_name(company_a, "㈱ サンプル建設") == customer

    def test_returns_none_when_not_registered(self, company_a):
        assert find_customer_by_name(company_a, "未登録工務店") is None

    def test_does_not_match_another_companys_customer(self, company_a, company_b):
        Customer.unscoped.create(
            company=company_b, code="C001", name="株式会社サンプル建設",
        )
        assert find_customer_by_name(company_a, "株式会社サンプル建設") is None

    def test_ignores_inactive_customer(self, company_a):
        Customer.unscoped.create(
            company=company_a, code="C002", name="休眠商事", is_active=False,
        )
        assert find_customer_by_name(company_a, "休眠商事") is None


# ---------------------------------------------------------------------------
# PDF 経路
# ---------------------------------------------------------------------------

class TestPdfPath:
    def test_pdf_uses_the_same_label_scan(self, tmp_path, monkeypatch):
        # pdfplumber の抽出結果は実ファイル依存なので、行の作りだけ差し替えて
        # 「PDF でも CSV と同じ走査に載る」ことを固定する。
        rows = [
            ["株式会社サンプル建設", "御中"],
            ["見積番号", "Q-2026-0142"],
            ["支払条件", "月末締め翌月末現金払い"],
            ["件名", "PDF から読んだ工事"],
        ]
        monkeypatch.setattr(importer, "_read_pdf_rows", lambda path: rows)

        path = tmp_path / "estimate.pdf"
        path.write_bytes(b"%PDF-1.4 dummy")
        data = parse_estimate_file(path, ".pdf")

        assert data["code"] == "Q-2026-0142"
        assert data["customer_name"] == "株式会社サンプル建設"
        assert data["payment_terms"] == "月末締め翌月末現金払い"
        assert data["name"] == "PDF から読んだ工事"


# ---------------------------------------------------------------------------
# 支払条件のフォールバック
# ---------------------------------------------------------------------------

class TestPaymentTermsFallback:
    def test_site_terms_win_over_customer_default(self, company_a):
        customer = Customer.unscoped.create(
            company=company_a, code="C001", name="得意先", payment_terms="現金",
        )
        site = Site.unscoped.create(
            company=company_a, code="S1", name="現場", customer=customer,
            payment_terms="手形120日",
        )
        assert site.effective_payment_terms == "手形120日"

    def test_falls_back_to_customer_default_when_blank(self, company_a):
        customer = Customer.unscoped.create(
            company=company_a, code="C001", name="得意先", payment_terms="現金",
        )
        site = Site.unscoped.create(
            company=company_a, code="S1", name="現場", customer=customer,
        )
        assert site.effective_payment_terms == "現金"

    def test_empty_when_neither_is_set(self, company_a):
        site = Site.unscoped.create(company=company_a, code="S1", name="現場")
        assert site.effective_payment_terms == ""


# ---------------------------------------------------------------------------
# 取り込み画面
# ---------------------------------------------------------------------------

@pytest.mark.django_db
class TestImportView:
    def _upload(self, client, tmp_path, csv_text=ESTIMATE_CSV):
        path = _write_csv(tmp_path, csv_text)
        with open(path, "rb") as f:
            return client.post(reverse("sites:import"), {"file": f})

    def test_upload_shows_the_parsed_values_for_confirmation(
        self, client, tmp_path, company_a, user_a
    ):
        set_current_company(company_a)
        client.force_login(user_a)

        res = self._upload(client, tmp_path)
        body = res.content.decode("utf-8")

        assert res.status_code == 200
        assert "Q-2026-0142" in body                    # 現場コードに見積番号
        assert "月末締め翌月末現金払い" in body            # 支払条件
        assert "○○ビル 電気設備改修工事" in body
        set_current_company(None)

    def test_warns_when_the_onchu_company_is_not_registered(
        self, client, tmp_path, company_a, user_a
    ):
        set_current_company(company_a)
        client.force_login(user_a)

        body = self._upload(client, tmp_path).content.decode("utf-8")

        assert "株式会社サンプル建設" in body
        assert "得意先マスタに登録されていません" in body
        set_current_company(None)

    def test_lists_fields_that_need_manual_entry(
        self, client, tmp_path, company_a, user_a
    ):
        set_current_company(company_a)
        client.force_login(user_a)

        body = self._upload(
            client, tmp_path, csv_text="件名,配線工事のみ\n"
        ).content.decode("utf-8")

        assert "手入力が必要な項目" in body
        assert "見積番号（現場コード）" in body
        assert "支払条件" in body
        set_current_company(None)

    def test_accepts_an_excel_file(self, client, tmp_path, company_a, user_a):
        set_current_company(company_a)
        client.force_login(user_a)

        path = _write_xlsx(tmp_path, {"表紙": ESTIMATE_SHEET})
        with open(path, "rb") as f:
            res = client.post(reverse("sites:import"), {"file": f})
        body = res.content.decode("utf-8")

        assert "Q-2026-0142" in body
        assert "月末締め翌月末現金払い" in body
        set_current_company(None)

    def test_rejects_an_unsupported_file_type(
        self, client, tmp_path, company_a, user_a
    ):
        set_current_company(company_a)
        client.force_login(user_a)

        path = tmp_path / "estimate.docx"
        path.write_bytes(b"dummy")
        with open(path, "rb") as f:
            res = client.post(reverse("sites:import"), {"file": f})

        assert "を選んでください。" in res.content.decode("utf-8")
        assert not Site.unscoped.filter(company=company_a).exists()
        set_current_company(None)

    def test_old_excel_format_is_rejected_with_guidance(
        self, client, tmp_path, company_a, user_a
    ):
        set_current_company(company_a)
        client.force_login(user_a)

        path = tmp_path / "estimate.xls"
        path.write_bytes(b"dummy")
        with open(path, "rb") as f:
            res = client.post(reverse("sites:import"), {"file": f})

        assert ".xlsx で保存し直して" in res.content.decode("utf-8")
        set_current_company(None)

    def test_confirming_creates_the_site_with_estimate_number_as_code(
        self, client, company_a, user_a
    ):
        set_current_company(company_a)
        customer = Customer.unscoped.create(
            company=company_a, code="C001", name="株式会社サンプル建設",
        )
        client.force_login(user_a)

        res = client.post(reverse("sites:import"), {
            "step": "confirm",
            "code": "Q-2026-0142",
            "name": "○○ビル 電気設備改修工事",
            "customer": customer.pk,
            "status": Site.Status.ESTIMATING,
            "contract_amount": 3480000,
            "payment_terms": "月末締め翌月末現金払い",
            "start_date": "2026-09-01",
            "end_date": "2026-11-30",
            "address": "山形県酒田市中町1-2-3",
        })

        site = Site.unscoped.get(company=company_a, code="Q-2026-0142")
        assert res.status_code == 302
        assert site.customer == customer
        assert site.payment_terms == "月末締め翌月末現金払い"
        assert site.company == company_a
        assert site.created_by == user_a
        set_current_company(None)

    def test_confirming_records_an_import_history_entry(
        self, client, company_a, user_a
    ):
        set_current_company(company_a)
        customer = Customer.unscoped.create(
            company=company_a, code="C001", name="株式会社サンプル建設",
        )
        client.force_login(user_a)

        client.post(reverse("sites:import"), {
            "step": "confirm",
            "filename": "見積書_本厚木.xlsx",
            "parsed_customer_name": "株式会社サンプル建設",
            "code": "Q-2026-0142",
            "name": "○○ビル 電気設備改修工事",
            "customer": customer.pk,
            "status": Site.Status.ESTIMATING,
            "contract_amount": 3480000,
            "payment_terms": "月末締め翌月末現金払い",
        })

        record = EstimateImport.unscoped.get(company=company_a)
        assert record.customer == customer
        assert record.customer_name_raw == "株式会社サンプル建設"
        assert record.filename == "見積書_本厚木.xlsx"
        assert record.estimate_number == "Q-2026-0142"
        assert record.amount == 3480000
        assert record.created_by == user_a
        assert record.site.code == "Q-2026-0142"
        assert record.is_unmatched is False
        set_current_company(None)

    def test_unmatched_company_name_is_kept_as_a_registration_candidate(
        self, client, company_a, user_a
    ):
        set_current_company(company_a)
        client.force_login(user_a)

        client.post(reverse("sites:import"), {
            "step": "confirm",
            "filename": "見積書.csv",
            "parsed_customer_name": "未登録工務店",
            "code": "Q-0002",
            "name": "宛名未登録の現場",
            "status": Site.Status.ESTIMATING,
            "contract_amount": 0,
        })

        record = EstimateImport.unscoped.get(company=company_a)
        assert record.customer is None
        assert record.customer_name_raw == "未登録工務店"
        assert record.is_unmatched is True
        set_current_company(None)

    def test_preview_alone_does_not_record_history(
        self, client, tmp_path, company_a, user_a
    ):
        # 読み取りを試しただけの操作は業務上の出来事ではないので残さない。
        set_current_company(company_a)
        client.force_login(user_a)

        self._upload(client, tmp_path)

        assert not EstimateImport.unscoped.filter(company=company_a).exists()
        set_current_company(None)

    def test_import_history_does_not_leak_across_tenants(
        self, client, company_a, company_b, user_a, user_b
    ):
        set_current_company(company_a)
        client.force_login(user_a)
        client.post(reverse("sites:import"), {
            "step": "confirm", "filename": "a.csv", "parsed_customer_name": "A社取引先",
            "code": "Q-A", "name": "A社の現場", "status": Site.Status.ESTIMATING,
            "contract_amount": 0,
        })
        set_current_company(None)

        set_current_company(company_b)
        assert EstimateImport.objects.count() == 0
        set_current_company(None)

    def test_other_details_are_saved_on_the_site(self, client, company_a, user_a):
        set_current_company(company_a)
        client.force_login(user_a)

        client.post(reverse("sites:import"), {
            "step": "confirm",
            "code": "Q-0003",
            "name": "その他項目つき",
            "status": Site.Status.ESTIMATING,
            "contract_amount": 0,
            "estimate_valid_until": "発行後30日間",
            "note": "足場は別途",
            "extracted_details": "工事区分: 電気設備工事\n現場代理人: 剣持",
        })

        site = Site.unscoped.get(company=company_a, code="Q-0003")
        assert site.estimate_valid_until == "発行後30日間"
        assert site.note == "足場は別途"
        assert "工事区分: 電気設備工事" in site.extracted_details
        set_current_company(None)

    def test_manual_entry_fills_in_what_the_file_did_not_have(
        self, client, company_a, user_a
    ):
        # ファイルに支払条件が無くても、確認画面で入力すれば保存される。
        set_current_company(company_a)
        client.force_login(user_a)

        client.post(reverse("sites:import"), {
            "step": "confirm",
            "code": "Q-0001",
            "name": "手入力した現場",
            "status": Site.Status.ESTIMATING,
            "contract_amount": 0,
            "payment_terms": "検収後60日",
        })

        site = Site.unscoped.get(company=company_a, code="Q-0001")
        assert site.payment_terms == "検収後60日"
        set_current_company(None)


# ---------------------------------------------------------------------------
# 取引先の詳細ページ（取込履歴の置き場所）
# ---------------------------------------------------------------------------

@pytest.mark.django_db
class TestCustomerDetail:
    def test_shows_the_import_history_for_that_customer(
        self, client, company_a, user_a
    ):
        set_current_company(company_a)
        customer = Customer.unscoped.create(
            company=company_a, code="C001", name="株式会社サンプル建設",
        )
        site = Site.unscoped.create(
            company=company_a, code="Q-1", name="取込で作った現場", customer=customer,
        )
        EstimateImport.unscoped.create(
            company=company_a, customer=customer, site=site,
            customer_name_raw="株式会社サンプル建設",
            filename="見積書_本厚木.xlsx", estimate_number="Q-1", amount=1000000,
        )
        client.force_login(user_a)

        body = client.get(
            reverse("masters:customer_detail", args=[customer.pk])
        ).content.decode("utf-8")

        assert "見積ファイルの取込履歴" in body
        assert "見積書_本厚木.xlsx" in body
        assert "取込で作った現場" in body     # 取引履歴（現場）にも出る
        set_current_company(None)

    def test_does_not_expose_another_companys_customer(
        self, client, company_a, company_b, user_a
    ):
        set_current_company(company_a)
        other = Customer.unscoped.create(
            company=company_b, code="C001", name="B社の得意先",
        )
        client.force_login(user_a)

        res = client.get(reverse("masters:customer_detail", args=[other.pk]))

        assert res.status_code == 404
        set_current_company(None)


# ---------------------------------------------------------------------------
# 既にある現場への取り込み
#
# 落札・受注のフェーズ移行で自動作成された現場は、件名と概算金額しか持たない。
# そこへ後からライデンの Excel を入れて数字を埋める経路。新規登録と違い
# 「項目ごとに反映するかを人が選ぶ」ことが要件なので、既定のチェック状態と、
# 選ばなかった項目が変わらないことを固定する。
# ---------------------------------------------------------------------------

@pytest.mark.django_db
class TestExistingSiteImport:
    def _site(self, company, **kwargs):
        fields = {
            "code": "BID-1",
            "name": "落札で自動作成された現場",
            "status": Site.Status.ORDERED,
            "contract_amount": 0,
        }
        fields.update(kwargs)
        return Site.unscoped.create(company=company, **fields)

    def _upload(self, client, tmp_path, site, sheets=None):
        path = _write_xlsx(tmp_path, sheets or {"表紙": ESTIMATE_SHEET})
        with open(path, "rb") as f:
            return client.post(
                reverse("sites:estimate_import", args=[site.pk]), {"file": f}
            )

    def test_upload_shows_current_and_file_values_side_by_side(
        self, client, tmp_path, company_a, user_a
    ):
        set_current_company(company_a)
        site = self._site(company_a)
        client.force_login(user_a)

        body = self._upload(client, tmp_path, site).content.decode("utf-8")

        assert "落札で自動作成された現場" in body       # 今の値
        assert "○○ビル 電気設備改修工事" in body       # ファイルの値
        assert "3480000" in body
        set_current_company(None)

    def test_empty_fields_are_checked_and_filled_fields_are_not(
        self, client, tmp_path, company_a, user_a
    ):
        # 空欄だけ既定でチェック。人が入れた現場名を黙って書き換えない。
        set_current_company(company_a)
        site = self._site(company_a)
        client.force_login(user_a)

        rows = {
            row["field"]: row
            for row in self._upload(client, tmp_path, site).context["diff_rows"]
        }

        assert rows["contract_amount"]["checked"] is True    # 0円なので空扱い
        assert rows["address"]["checked"] is True
        assert rows["start_date"]["checked"] is True
        assert rows["name"]["checked"] is False              # 既に件名がある
        assert rows["name"]["current"] == "落札で自動作成された現場"
        set_current_company(None)

    def test_only_the_checked_fields_are_applied(self, client, company_a, user_a):
        set_current_company(company_a)
        site = self._site(company_a)
        client.force_login(user_a)

        res = client.post(reverse("sites:estimate_import", args=[site.pk]), {
            "step": "confirm",
            "filename": "ライデン見積.xlsx",
            "apply": ["contract_amount", "start_date", "end_date", "address"],
            "value_contract_amount": "3480000",
            "value_start_date": "2026-09-01",
            "value_end_date": "2026-11-30",
            "value_address": "山形県酒田市中町1-2-3",
            "value_name": "○○ビル 電気設備改修工事",
            "value_payment_terms": "月末締め翌月末現金払い",
        })

        site.refresh_from_db()
        assert res.status_code == 302
        assert site.contract_amount == 3480000
        assert site.start_date == datetime.date(2026, 9, 1)
        assert site.end_date == datetime.date(2026, 11, 30)
        assert site.address == "山形県酒田市中町1-2-3"
        # チェックを外した項目は今の値のまま
        assert site.name == "落札で自動作成された現場"
        assert site.payment_terms == ""
        set_current_company(None)

    def test_a_too_long_value_is_truncated_instead_of_failing_to_save(
        self, client, company_a, user_a
    ):
        # 明細行を掴む等で長い値が来ても、保存時に落とさず切って入れる。
        set_current_company(company_a)
        site = self._site(company_a, code="")
        client.force_login(user_a)

        client.post(reverse("sites:estimate_import", args=[site.pk]), {
            "step": "confirm",
            "apply": ["code"],
            "value_code": "Q" * 80,
        })

        site.refresh_from_db()
        assert len(site.code) == 50
        set_current_company(None)

    def test_note_is_appended_so_site_side_remarks_survive(
        self, client, company_a, user_a
    ):
        set_current_company(company_a)
        site = self._site(company_a, note="鍵は警備室で受け取ること")
        client.force_login(user_a)

        client.post(reverse("sites:estimate_import", args=[site.pk]), {
            "step": "confirm",
            "apply": ["note"],
            "value_note": "足場は別途",
        })

        site.refresh_from_db()
        assert "鍵は警備室で受け取ること" in site.note
        assert "足場は別途" in site.note
        set_current_company(None)

    def test_importing_the_same_file_twice_does_not_duplicate_lines(
        self, client, company_a, user_a
    ):
        set_current_company(company_a)
        site = self._site(company_a)
        client.force_login(user_a)
        payload = {
            "step": "confirm",
            "apply": ["extracted_details"],
            "value_extracted_details": "工事区分: 電気設備工事",
        }

        client.post(reverse("sites:estimate_import", args=[site.pk]), payload)
        client.post(reverse("sites:estimate_import", args=[site.pk]), payload)

        site.refresh_from_db()
        assert site.extracted_details.count("工事区分: 電気設備工事") == 1
        set_current_company(None)

    def test_applying_records_an_import_history_entry_on_the_site(
        self, client, company_a, user_a
    ):
        set_current_company(company_a)
        site = self._site(company_a)
        client.force_login(user_a)

        client.post(reverse("sites:estimate_import", args=[site.pk]), {
            "step": "confirm",
            "filename": "ライデン見積.xlsx",
            "parsed_customer_name": "株式会社サンプル建設",
            "apply": ["contract_amount"],
            "value_contract_amount": "3480000",
        })

        record = EstimateImport.unscoped.get(company=company_a)
        assert record.site == site
        assert record.filename == "ライデン見積.xlsx"
        assert record.customer_name_raw == "株式会社サンプル建設"
        assert record.amount == 3480000
        assert record.created_by == user_a
        set_current_company(None)

    def test_selecting_nothing_changes_nothing_and_leaves_no_history(
        self, client, company_a, user_a
    ):
        set_current_company(company_a)
        site = self._site(company_a)
        client.force_login(user_a)

        client.post(reverse("sites:estimate_import", args=[site.pk]), {
            "step": "confirm",
            "filename": "ライデン見積.xlsx",
            "value_contract_amount": "3480000",
        })

        site.refresh_from_db()
        assert site.contract_amount == 0
        assert not EstimateImport.unscoped.filter(company=company_a).exists()
        set_current_company(None)

    def test_preview_alone_does_not_change_the_site(
        self, client, tmp_path, company_a, user_a
    ):
        set_current_company(company_a)
        site = self._site(company_a)
        client.force_login(user_a)

        self._upload(client, tmp_path, site)

        site.refresh_from_db()
        assert site.contract_amount == 0
        assert not EstimateImport.unscoped.filter(company=company_a).exists()
        set_current_company(None)

    def test_matched_customer_can_be_applied_to_the_site(
        self, client, tmp_path, company_a, user_a
    ):
        set_current_company(company_a)
        customer = Customer.unscoped.create(
            company=company_a, code="C001", name="㈱サンプル建設",
        )
        site = self._site(company_a)
        client.force_login(user_a)

        self._upload(client, tmp_path, site)
        client.post(reverse("sites:estimate_import", args=[site.pk]), {
            "step": "confirm",
            "apply": ["customer"],
            "value_customer": str(customer.pk),
        })

        site.refresh_from_db()
        assert site.customer == customer
        set_current_company(None)

    def test_another_companys_customer_cannot_be_applied(
        self, client, company_a, company_b, user_a
    ):
        # value_customer は画面から往復する値なので、他社の pk は弾く。
        set_current_company(company_a)
        other = Customer.unscoped.create(
            company=company_b, code="C001", name="B社の得意先",
        )
        site = self._site(company_a)
        client.force_login(user_a)

        client.post(reverse("sites:estimate_import", args=[site.pk]), {
            "step": "confirm",
            "apply": ["customer"],
            "value_customer": str(other.pk),
        })

        site.refresh_from_db()
        assert site.customer is None
        set_current_company(None)

    def test_another_companys_site_is_not_reachable(
        self, client, company_a, company_b, user_a
    ):
        set_current_company(company_a)
        other_site = self._site(company_b, code="BID-B")
        client.force_login(user_a)

        res = client.get(reverse("sites:estimate_import", args=[other_site.pk]))

        assert res.status_code == 404
        set_current_company(None)
