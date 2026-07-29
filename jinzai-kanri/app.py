"""人材管理システム — Streamlit メインアプリ。

社員の追加・編集・削除・一覧表示・Excel出力を行う。
master.employees テーブルを直接管理する。
"""
from __future__ import annotations

import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
import pandas as pd
import streamlit as st
import streamlit.components.v1 as _components

import config
import database as db

st.set_page_config(page_title="人材管理", page_icon="👥", layout="wide")


def _disable_browser_translation():
    _components.html(
        """<script>
        try {
          const doc = window.parent.document;
          doc.documentElement.setAttribute('translate', 'no');
          doc.documentElement.lang = 'ja';
          doc.documentElement.classList.add('notranslate');
          if (!doc.querySelector('meta[name="google"][content="notranslate"]')) {
            const m = doc.createElement('meta');
            m.name = 'google'; m.content = 'notranslate';
            doc.head.appendChild(m);
          }
        } catch (e) {}
        </script>""",
        height=0,
    )


_disable_browser_translation()
db.init_db()


# ========================================================================
# 画面: 社員一覧
# ========================================================================
def page_list() -> None:
    st.header("👥 社員一覧")

    # フィルタ
    c1, c2, c3 = st.columns(3)
    role_opts = ["全員"] + config.ROLES
    selected_role = c1.radio("職種区分", role_opts, horizontal=True, key="lst_role")
    show_inactive = c2.checkbox("退職者も表示", key="lst_inactive")
    keyword = c3.text_input("名前検索", placeholder="氏名・よみがなで検索", key="lst_kw")

    role_filter = None if selected_role == "全員" else selected_role
    active_only = not show_inactive
    employees = db.list_employees(active_only=active_only, role=role_filter)

    if keyword.strip():
        kw = keyword.strip().lower()
        employees = [e for e in employees
                     if kw in (e["name"] or "").lower()
                     or kw in (e.get("kana") or "").lower()]

    if not employees:
        st.info("該当する社員がいません。")
        return

    st.caption(f"表示中: {len(employees)} 名（50音順）")

    # テーブル表示
    df = pd.DataFrame(employees)
    df["在籍"] = df["is_active"].map({True: "○", False: "✕"})
    view_cols = {
        "code": "社員番号", "name": "氏名", "kana": "よみがな",
        "department": "部署", "position": "役職", "role": "職種区分",
        "phone": "電話番号", "在籍": "在籍", "note": "備考",
    }
    st.dataframe(
        df[list(view_cols.keys())].rename(columns=view_cols),
        use_container_width=True, hide_index=True,
    )

    # Excel 出力
    st.divider()
    _excel_download(employees)


# ========================================================================
# 画面: 社員登録
# ========================================================================
def page_add() -> None:
    st.header("➕ 社員登録")

    if st.session_state.get("add_flash"):
        st.success(st.session_state.pop("add_flash"))

    with st.form("add_employee", clear_on_submit=True):
        c1, c2 = st.columns(2)
        code = c1.text_input("社員番号", placeholder="E011")
        name = c2.text_input("氏名（漢字）", placeholder="田中 太郎")
        c3, c4 = st.columns(2)
        kana = c3.text_input("よみがな", placeholder="たなか たろう")
        phone = c4.text_input("電話番号", placeholder="090-1234-5678")
        c5, c6, c7 = st.columns(3)
        department = c5.selectbox("部署", [""] + config.DEPARTMENTS)
        position = c6.selectbox("役職", [""] + config.POSITIONS)
        role = c7.selectbox("職種区分", config.ROLES)
        note = st.text_area("備考（資格・特記事項）", height=80, placeholder="第一種電気工事士 等")

        if st.form_submit_button("登録", type="primary"):
            if not name.strip():
                st.error("氏名を入力してください。")
            else:
                emp_id = db.add_employee(
                    code=code.strip(), name=name.strip(), kana=kana.strip(),
                    department=department, position=position, role=role,
                    phone=phone.strip(), note=note.strip(),
                )
                st.session_state["add_flash"] = f"{name.strip()} を登録しました（ID: {emp_id}）。"
                st.rerun()


# ========================================================================
# 画面: 社員編集
# ========================================================================
def page_edit() -> None:
    st.header("✏️ 社員編集")

    employees = db.list_employees()
    if not employees:
        st.info("社員が登録されていません。")
        return

    opts = {f"#{e['id']} {e['code']} {e['name']}": e for e in employees}
    pick = st.selectbox("編集する社員", list(opts.keys()), key="edit_pick")
    emp = opts[pick]

    if st.session_state.get("edit_flash"):
        st.success(st.session_state.pop("edit_flash"))

    c1, c2 = st.columns(2)
    new_code = c1.text_input("社員番号", value=emp["code"] or "", key=f"ed_code_{emp['id']}")
    new_name = c2.text_input("氏名", value=emp["name"], key=f"ed_name_{emp['id']}")
    c3, c4 = st.columns(2)
    new_kana = c3.text_input("よみがな", value=emp.get("kana") or "", key=f"ed_kana_{emp['id']}")
    new_phone = c4.text_input("電話番号", value=emp.get("phone") or "", key=f"ed_phone_{emp['id']}")
    c5, c6, c7 = st.columns(3)
    dept_opts = [""] + config.DEPARTMENTS
    dept_idx = dept_opts.index(emp.get("department") or "") if (emp.get("department") or "") in dept_opts else 0
    new_dept = c5.selectbox("部署", dept_opts, index=dept_idx, key=f"ed_dept_{emp['id']}")
    pos_opts = [""] + config.POSITIONS
    pos_idx = pos_opts.index(emp.get("position") or "") if (emp.get("position") or "") in pos_opts else 0
    new_pos = c6.selectbox("役職", pos_opts, index=pos_idx, key=f"ed_pos_{emp['id']}")
    role_opts = [""] + config.ROLES
    cur_role = emp.get("role") or ""
    role_idx = role_opts.index(cur_role) if cur_role in role_opts else 0
    new_role = c7.selectbox("職種区分", role_opts, index=role_idx, key=f"ed_role_{emp['id']}")
    new_note = st.text_area("備考", value=emp.get("note") or "", height=80, key=f"ed_note_{emp['id']}")
    new_active = st.selectbox("在籍状態", ["在籍", "退職"],
                              index=0 if emp["is_active"] else 1, key=f"ed_active_{emp['id']}")

    bc1, bc2 = st.columns([3, 1])
    if bc1.button("変更を保存", type="primary", use_container_width=True, key="ed_save"):
        if not new_name.strip():
            st.error("氏名を入力してください。")
        else:
            db.update_employee(
                emp["id"],
                code=new_code.strip(), name=new_name.strip(), kana=new_kana.strip(),
                department=new_dept, position=new_pos, role=new_role,
                phone=new_phone.strip(), note=new_note.strip(),
                is_active=new_active == "在籍",
            )
            st.session_state["edit_flash"] = f"{new_name.strip()} の情報を更新しました。"
            st.rerun()

    # 削除
    with st.expander("🗑️ この社員を削除"):
        st.warning("削除すると、この社員のマスター情報が完全に消えます。通常は「退職」にしてください。")
        if st.button("完全に削除する", key="ed_del"):
            db.delete_employee(emp["id"])
            st.success(f"{emp['name']} を削除しました。")
            st.rerun()


# ========================================================================
# Excel出力
# ========================================================================
def _excel_download(employees: list[dict]) -> None:
    """社員一覧をExcelファイルとしてダウンロード。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "社員名簿"

    headers = [
        ("社員番号", 12), ("氏名", 18), ("よみがな", 22),
        ("部署", 16), ("役職", 16), ("職種区分", 14),
        ("電話番号", 18), ("在籍", 10), ("備考", 40),
    ]
    fields = ["code", "name", "kana", "department", "position", "role", "phone", "is_active", "note"]

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    even_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")

    for col, (label, width) in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=label)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    for i, emp in enumerate(employees):
        row = i + 2
        for col, field in enumerate(fields, 1):
            val = emp.get(field, "")
            if field == "is_active":
                val = "在籍" if val else "退職"
            c = ws.cell(row=row, column=col, value=val or "")
            c.border = thin_border
            if i % 2 == 1:
                c.fill = even_fill

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    now_str = datetime.now().strftime("%Y%m%d")
    st.download_button(
        label="📥 Excelダウンロード",
        data=buf.getvalue(),
        file_name=f"社員名簿_{now_str}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )


# ========================================================================
# JSON同期
# ========================================================================
def page_sync() -> None:
    st.header("🔄 データ同期")
    st.caption("DBの社員マスターをリポジトリのJSONファイルに書き出し/読み込みします。")

    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**DB → JSON に書き出し**")
        st.caption(f"保存先: `{config.EMPLOYEES_JSON}`")
        if st.button("DBの内容をJSONに保存", use_container_width=True):
            db.save_to_json()
            st.success("JSONファイルを更新しました。")

    with c2:
        st.markdown("**JSON → DB に読み込み**")
        st.caption("DBの全社員を削除してJSONから再投入します。")
        if st.button("JSONからDBを再構築", type="secondary", use_container_width=True):
            with get_conn_for_reset() as conn:
                with conn.cursor() as cur:
                    cur.execute("DELETE FROM master.employees")
            db._restore_if_empty()
            st.success("JSONからDBを再構築しました。")
            st.rerun()


def get_conn_for_reset():
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "shared"))
    from dbconn import get_conn
    return get_conn()


# ========================================================================
# ルーティング
# ========================================================================
PAGES = {
    "社員一覧": page_list,
    "社員登録": page_add,
    "社員編集": page_edit,
    "データ同期": page_sync,
}

st.sidebar.title("👥 人材管理")
st.sidebar.caption(config.COMPANY_NAME)

if "page" not in st.session_state:
    st.session_state["page"] = "社員一覧"
choice = st.sidebar.radio("メニュー", list(PAGES.keys()),
                          index=list(PAGES.keys()).index(st.session_state["page"]))
if choice != st.session_state["page"]:
    st.session_state["page"] = choice

st.sidebar.divider()
emp_count = len(db.list_employees(active_only=True))
st.sidebar.metric("在籍社員数", f"{emp_count} 名")

PAGES[st.session_state["page"]]()
